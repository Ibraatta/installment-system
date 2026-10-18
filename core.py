"""Excel is the system of record. All cash values are integer halalas in the engine."""
from __future__ import annotations
from calendar import monthrange
from datetime import date, datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo
import json
import hashlib
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

VERSION = '1.0.0'
TZ = ZoneInfo('Asia/Riyadh')
SOURCE = 'https://fabulous-lipstick-e9d.notion.site/39d57dc9757a8008a982c7af38a4c0de'
SCHEMA = {
 'العملاء': ['رقم العميل','اسم العميل','رقم الهوية','رقم الجوال','تاريخ إضافة العميل','حالة العميل','التصنيف الائتماني','الحد الائتماني','العنوان','جهة العمل','الدخل الشهري','اسم الكفيل','جوال الكفيل','ملاحظات','إجمالي التمويل','إجمالي المتبقي','المتبقي من التمويل','عدد مرات التأخير','إجمالي العقود','آخر تعديل'],
 'العقود': ['رقم العقد','رقم العميل','تاريخ العقد','مبلغ التمويل','إجمالي السداد','مدة العقد','أول قسط','رسوم العقد','ملاحظات','حالة الصرف','مبلغ القسط','إجمالي المحصل','المبلغ المتبقي','الربح المتوقع','عدد الأقساط المسددة','عدد الأقساط المتبقية','المتبقي من التمويل','الاستحقاق القادم','تاريخ انتهاء العقد','حالة العقد','خصم السداد المبكر','آخر تعديل'],
 'الأقساط': ['رقم القسط','رقم العقد','تسلسل القسط','تاريخ الاستحقاق','مبلغ القسط','المحصل','المتبقي','حالة القسط','أيام التأخير','تاريخ اكتمال السداد','خصم التسوية'],
 'الحركات المالية': ['رقم العملية','التاريخ','المبلغ','نوع العملية','رقم العقد','المحصل','المرجع','ملاحظات','أثر الحركة','خصم التسوية','حالة العملية','وقت التسجيل','رقم الطلب','إلغاء بواسطة'],
 'توزيع التحصيل': ['رقم التوزيع','رقم العملية','رقم القسط','المبلغ','نوع التوزيع','تاريخ العملية'],
 'متابعة العملاء': ['رقم المتابعة','رقم العميل','رقم العقد','تاريخ التواصل','وسيلة التواصل','نتيجة التواصل','وعد السداد','المبلغ الموعود','المتابعة القادمة','الحالة','ملاحظات','آخر تعديل'],
 'الإعدادات': ['المفتاح','القيمة','الوصف'],
 'سجل التغييرات': ['رقم السجل','الوقت','نوع الإجراء','رقم السجل المتأثر','تفاصيل'],
}
CASH_TYPES = {'رأس مال':1,'إضافة رأس مال':1,'سحب رأس مال':-1,'صرف تمويل':-1,'سداد قسط':1,'سداد مبكر':1,'رسوم عقد':1,'مصروف':-1}
RECEIPTS = {'سداد قسط','سداد مبكر'}
CLIENT_STATES = ['نشط','موقوف','مؤرشف']
GRADES = ['غير مصنف','A','B','C','D']
class ValidationError(ValueError): pass
class ConflictError(RuntimeError): pass

def today(): return datetime.now(TZ).date()
def stamp(): return datetime.now(TZ).isoformat(timespec='microseconds')
def uid(prefix): return f'{prefix}-{uuid4().hex[:12].upper()}'
def money(value):
    try:
        x = Decimal(str(value or 0))
        if not x.is_finite(): raise ValueError()
        return int((x*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    except Exception: raise ValidationError('المبلغ غير صالح.')
def sar(value): return float(Decimal(value)/100)
def day(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    if v: return date.fromisoformat(str(v)[:10])
    return None

def text(v):
    s = str(v or '').strip()
    if len(s)>4000: raise ValidationError('النص أطول من الحد المسموح.')
    # Store user text as an Excel string, never as an executable formula.
    return s

def put(ws,row,col,value):
    c = ws.cell(row,col)
    c.value=value
    if isinstance(value,str): c.data_type='s'

def rows(wb,sheet):
    headers=SCHEMA[sheet]
    return [dict(zip(headers,r[:len(headers)])) for r in wb[sheet].iter_rows(min_row=2,values_only=True) if r[0]]
def append(wb,sheet,record):
    ws=wb[sheet]; r=ws.max_row+1
    for c,k in enumerate(SCHEMA[sheet],1): put(ws,r,c,record.get(k))
def update(wb,sheet,key,value,changes):
    ws=wb[sheet]; h=SCHEMA[sheet]; col=h.index(key)+1
    for r in range(2,ws.max_row+1):
        if ws.cell(r,col).value==value:
            for k,v in changes.items(): put(ws,r,h.index(k)+1,v)
            return
    raise ValidationError('السجل غير موجود.')
def find(wb,sheet,key,value):
    return next((r for r in rows(wb,sheet) if r[key]==value),None)
def require(wb,sheet,key,value):
    r=find(wb,sheet,key,value)
    if not r: raise ValidationError('السجل المرتبط غير موجود.')
    return r

def settings(wb): return {r['المفتاح']:r['القيمة'] for r in rows(wb,'الإعدادات')}
def add_month(anchor,n):
    y,m=divmod(anchor.year*12+anchor.month-1+n,12); m+=1
    return date(y,m,min(anchor.day,monthrange(y,m)[1]))
def active_tx(wb): return [r for r in rows(wb,'الحركات المالية') if r['حالة العملية']=='نشطة']
def cash(wb): return sum(money(r['المبلغ'])*CASH_TYPES[r['نوع العملية']] for r in active_tx(wb))

def initialize(capital=100000,capital_date=date(2026,7,1)):
    wb=Workbook(); wb.remove(wb.active)
    wb.create_sheet('ملخص المحفظة')
    for name,h in SCHEMA.items(): wb.create_sheet(name).append(h)
    for key,val,desc in [
       ('schema_version',1,'بنية الملف؛ لا تعدل هذا السطر'),
       ('default_months',16,'مدة افتراضية قابلة للتعديل لكل عقد'),
       ('default_profit_rate',0.60,'نسبة ربح إجمالية افتراضية قابلة للتعديل؛ ليست نسبة شهرية'),
       ('currency','SAR','الريال السعودي'),
       ('source',SOURCE,'الصفحة المرجعية للتقسيم والحقول'),
       ('principal_method','proportional','المتبقي من التمويل تقدير نسبي وليس قاعدة تسوية مبكرة'),
       ('calculation_note','إجمالي السداد يحدده المستخدم. توزيع السداد للأقدم، وأي فرق هللات في القسط الأخير.','صيغ نوشن الداخلية غير متاحة من رابط العرض العام'),
       ('grading_note','تصنيف داخلي يدوي؛ لم ينقل التطبيق معادلة تصنيف غير ظاهرة من نوشن.','التأخر يظهر من سجلات الاستحقاق والسداد'),
    ]: append(wb,'الإعدادات',dict(zip(SCHEMA['الإعدادات'],[key,val,desc])))
    if money(capital):
        tx(wb,'رأس مال',capital,capital_date,notes='رصيد البداية الظاهر في صفحة نوشن المرجعية',request_id='NOTION-OPENING-20260701')
    rebuild(wb); format_book(wb)
    return wb

def tx(wb,kind,amount,when,contract='',collector='',reference='',notes='',discount=0,request_id=None):
    if kind not in CASH_TYPES: raise ValidationError('نوع الحركة غير صالح.')
    amount_h=money(amount)
    if amount_h<=0: raise ValidationError('المبلغ يجب أن يكون أكبر من صفر.')
    when=day(when)
    if not when or when>today(): raise ValidationError('اختر تاريخ حركة صحيحًا لا يتجاوز اليوم.')
    request_id=request_id or uid('REQ')
    # A double click or a lost upload response must never duplicate a payment.
    if any(r['رقم الطلب']==request_id for r in rows(wb,'الحركات المالية')):
        return next(r['رقم العملية'] for r in rows(wb,'الحركات المالية') if r['رقم الطلب']==request_id)
    reference=text(reference)
    if reference and any(r['المرجع']==reference and r['حالة العملية']=='نشطة' for r in rows(wb,'الحركات المالية')):
        raise ValidationError('المرجع مستخدم في حركة نشطة؛ راجع احتمال تكرار التسجيل.')
    cid=text(contract)
    c=None
    if kind in RECEIPTS|{'صرف تمويل','رسوم عقد'}:
        c=require(wb,'العقود','رقم العقد',cid)
        if when<day(c['تاريخ العقد']): raise ValidationError('الحركة لا يمكن أن تسبق تاريخ العقد.')
    elif cid: raise ValidationError('هذا النوع من الحركة لا يرتبط بعقد.')
    if kind=='صرف تمويل':
        client=require(wb,'العملاء','رقم العميل',c['رقم العميل'])
        if client['حالة العميل']!='نشط': raise ValidationError('العميل غير نشط.')
        if any(r['رقم العقد']==cid and r['نوع العملية']=='صرف تمويل' for r in active_tx(wb)):
            raise ValidationError('سبق صرف هذا العقد.')
        if amount_h!=money(c['مبلغ التمويل']): raise ValidationError('مبلغ الصرف يجب أن يساوي مبلغ تمويل العقد.')
        snap=snapshot(wb)
        outstanding=sum(x['principal_remaining_h'] for x in snap['contracts'] if x['client_id']==c['رقم العميل'] and x['disbursed'])
        limit=money(client['الحد الائتماني'])
        if limit and outstanding+amount_h>limit: raise ValidationError('الصرف يتجاوز الحد الائتماني المحدد للعميل.')
    ident=uid('TX')
    append(wb,'الحركات المالية',{
        'رقم العملية':ident,'التاريخ':when,'المبلغ':sar(amount_h),'نوع العملية':kind,'رقم العقد':cid,
        'المحصل':text(collector),'المرجع':reference,'ملاحظات':text(notes),'أثر الحركة':sar(amount_h*CASH_TYPES[kind]),
        'خصم التسوية':sar(money(discount)),'حالة العملية':'نشطة','وقت التسجيل':stamp(),'رقم الطلب':request_id,
    })
    return ident

def add_client(wb,data):
    name=text(data.get('اسم العميل')); identity=text(data.get('رقم الهوية')); phone=text(data.get('رقم الجوال'))
    if not name or not phone: raise ValidationError('اسم العميل ورقم الجوال مطلوبان.')
    if identity and any(r['رقم الهوية']==identity for r in rows(wb,'العملاء')): raise ValidationError('رقم الهوية مسجل لعميل آخر.')
    if data.get('حالة العميل','نشط') not in CLIENT_STATES or data.get('التصنيف الائتماني','غير مصنف') not in GRADES:
        raise ValidationError('حالة أو تصنيف غير صالح.')
    for k in ['الحد الائتماني','الدخل الشهري']:
        if money(data.get(k))<0: raise ValidationError('القيم المالية لا تقبل السالب.')
    record={k:text(data.get(k)) for k in SCHEMA['العملاء'][:14]}
    record.update({'رقم العميل':uid('CL'),'اسم العميل':name,'رقم الهوية':identity,'رقم الجوال':phone,
        'تاريخ إضافة العميل':day(data.get('تاريخ إضافة العميل')) or today(),'حالة العميل':data.get('حالة العميل','نشط'),
        'التصنيف الائتماني':data.get('التصنيف الائتماني','غير مصنف'),'الحد الائتماني':sar(money(data.get('الحد الائتماني'))),
        'الدخل الشهري':sar(money(data.get('الدخل الشهري'))),'آخر تعديل':stamp()})
    append(wb,'العملاء',record); return record['رقم العميل']

def edit_client(wb,ident,data):
    old=require(wb,'العملاء','رقم العميل',ident)
    identity=text(data.get('رقم الهوية'))
    if identity and any(r['رقم العميل']!=ident and r['رقم الهوية']==identity for r in rows(wb,'العملاء')):
        raise ValidationError('رقم الهوية مكرر.')
    if not text(data.get('اسم العميل')) or not text(data.get('رقم الجوال')): raise ValidationError('اسم العميل والجوال مطلوبان.')
    if data.get('حالة العميل') not in CLIENT_STATES or data.get('التصنيف الائتماني') not in GRADES: raise ValidationError('تصنيف غير صالح.')
    if data.get('حالة العميل')=='مؤرشف' and any(x['client_id']==ident and x['disbursed'] and x['remaining_h']>0 for x in snapshot(wb)['contracts']):
        raise ValidationError('لا يمكن أرشفة عميل لديه عقد غير مسدد.')
    changes={k:text(data.get(k,old.get(k))) for k in SCHEMA['العملاء'][1:14] if k!='تاريخ إضافة العميل'}
    for k in ['الحد الائتماني','الدخل الشهري']:
        if money(data.get(k))<0: raise ValidationError('المبلغ لا يقبل السالب.')
        changes[k]=sar(money(data.get(k)))
    changes['آخر تعديل']=stamp(); update(wb,'العملاء','رقم العميل',ident,changes)

def add_contract(wb,data,ident=None):
    client=require(wb,'العملاء','رقم العميل',data['رقم العميل'])
    if client['حالة العميل']!='نشط': raise ValidationError('يجب أن يكون العميل نشطًا.')
    p,t=money(data['مبلغ التمويل']),money(data['إجمالي السداد'])
    n=int(data['مدة العقد']); start=day(data['أول قسط']); signed=day(data['تاريخ العقد'])
    if p<=0 or t<p or n<1 or n>120 or t<n: raise ValidationError('تحقق من مبلغ التمويل وإجمالي السداد والمدة (1–120 شهرًا).')
    if not start or not signed or start<signed: raise ValidationError('تاريخ أول قسط لا يمكن أن يسبق العقد.')
    if signed>today(): raise ValidationError('تاريخ العقد لا يمكن أن يتجاوز اليوم.')
    fees=money(data.get('رسوم العقد'))
    if fees<0: raise ValidationError('رسوم العقد لا تقبل السالب.')
    ident=ident or text(data.get('رقم العقد')) or uid('CT')
    if find(wb,'العقود','رقم العقد',ident): raise ValidationError('رقم العقد مستخدم بالفعل.')
    append(wb,'العقود',{'رقم العقد':ident,'رقم العميل':data['رقم العميل'],'تاريخ العقد':signed,'مبلغ التمويل':sar(p),
       'إجمالي السداد':sar(t),'مدة العقد':n,'أول قسط':start,'رسوم العقد':sar(fees),'ملاحظات':text(data.get('ملاحظات')),'آخر تعديل':stamp()})
    base=t//n
    for i in range(n):
        append(wb,'الأقساط',{'رقم القسط':f'{ident}-{i+1:03}','رقم العقد':ident,'تسلسل القسط':i+1,
           'تاريخ الاستحقاق':add_month(start,i),'مبلغ القسط':sar(base if i<n-1 else t-base*(n-1))})
    return ident

def edit_contract(wb,ident,data):
    require(wb,'العقود','رقم العقد',ident)
    if any(r['رقم العقد']==ident for r in active_tx(wb)): raise ValidationError('تعديل الشروط متاح للمسودة التي لا ترتبط بحركات مالية فقط.')
    for sheet in ['العقود','الأقساط']:
        existing=[r for r in rows(wb,sheet) if r['رقم العقد']!=ident]
        ws=wb[sheet]; ws.delete_rows(2,ws.max_row)
        for r in existing: append(wb,sheet,r)
    add_contract(wb,data,ident)

def add_followup(wb,data):
    require(wb,'العملاء','رقم العميل',data['رقم العميل'])
    cid=data.get('رقم العقد','')
    if cid and require(wb,'العقود','رقم العقد',cid)['رقم العميل']!=data['رقم العميل']:
        raise ValidationError('العقد لا يخص العميل المختار.')
    if money(data.get('المبلغ الموعود'))<0: raise ValidationError('مبلغ الوعد لا يقبل السالب.')
    r={k:text(data.get(k)) for k in SCHEMA['متابعة العملاء']}
    for k in ['تاريخ التواصل','وعد السداد','المتابعة القادمة']: r[k]=day(data.get(k))
    r.update({'رقم المتابعة':uid('FU'),'المبلغ الموعود':sar(money(data.get('المبلغ الموعود'))),'آخر تعديل':stamp(),'الحالة':'مفتوحة'})
    append(wb,'متابعة العملاء',r); return r['رقم المتابعة']

def rebuild(wb):
    """Replay the active ledger and allocate receipts FIFO with exact cent rounding."""
    clients={r['رقم العميل']:r for r in rows(wb,'العملاء')}
    contracts={r['رقم العقد']:r for r in rows(wb,'العقود')}
    schedule=rows(wb,'الأقساط')
    installments={r['رقم القسط']:{**r,'paid_h':0,'waived_h':0,'paid_date':None} for r in schedule}
    if len(installments)!=len(schedule): raise ValidationError('رقم قسط مكرر.')
    for cid,c in contracts.items():
        if c['رقم العميل'] not in clients: raise ValidationError('عقد مرتبط بعميل غير موجود.')
        parts=[r for r in schedule if r['رقم العقد']==cid]
        if len(parts)!=int(c['مدة العقد']) or sum(money(x['مبلغ القسط']) for x in parts)!=money(c['إجمالي السداد']):
            raise ValidationError('جدول الأقساط لا يطابق مدة العقد أو إجمالي السداد.')
    alltx=active_tx(wb)
    if len({r['رقم العملية'] for r in rows(wb,'الحركات المالية')})!=len(rows(wb,'الحركات المالية')):
        raise ValidationError('أرقام حركات مكررة.')
    refs=[r['المرجع'] for r in alltx if r['المرجع']]
    if len(refs)!=len(set(refs)): raise ValidationError('مراجع مالية مكررة.')
    allocated=[]; running=0; disbursed={}; settled=set(); fee_paid={}
    for tr in sorted(alltx,key=lambda r:(day(r['التاريخ']),str(r['وقت التسجيل']),r['رقم العملية'])):
        typ=tr['نوع العملية']; amount=money(tr['المبلغ']); cid=tr['رقم العقد']; d=day(tr['التاريخ'])
        if typ not in CASH_TYPES or amount<=0 or d is None or d>today(): raise ValidationError('حركة مالية غير صالحة.')
        if typ in RECEIPTS|{'صرف تمويل','رسوم عقد'}:
            if cid not in contracts: raise ValidationError('حركة مرتبطة بعقد غير موجود.')
            if d<day(contracts[cid]['تاريخ العقد']): raise ValidationError('الحركة تسبق تاريخ العقد.')
        if typ=='صرف تمويل':
            if cid in disbursed: raise ValidationError('صرف مكرر لنفس العقد.')
            if amount!=money(contracts[cid]['مبلغ التمويل']): raise ValidationError('مبلغ الصرف يخالف مبلغ العقد.')
            disbursed[cid]=d
        if typ=='رسوم عقد':
            if cid not in disbursed: raise ValidationError('تحصيل الرسوم يتطلب صرف التمويل أولًا.')
            fee_paid[cid]=fee_paid.get(cid,0)+amount
            if fee_paid[cid]>money(contracts[cid]['رسوم العقد']): raise ValidationError('الرسوم المحصلة تتجاوز رسوم العقد.')
        if typ in RECEIPTS:
            if cid not in disbursed: raise ValidationError('لا يمكن تسجيل سداد لعقد لم يصرف أو بتاريخ يسبق الصرف.')
            if cid in settled: raise ValidationError('العقد مسدد مبكرًا؛ ألغِ التسوية قبل إضافة سداد آخر.')
            parts=sorted([r for r in installments.values() if r['رقم العقد']==cid],key=lambda x:x['تسلسل القسط'])
            remaining=sum(money(r['مبلغ القسط'])-r['paid_h']-r['waived_h'] for r in parts)
            waive=money(tr['خصم التسوية'])
            if amount>remaining or waive<0 or (typ=='سداد قسط' and waive): raise ValidationError('السداد أو خصم التسوية غير صالح.')
            if typ=='سداد مبكر' and amount+waive!=remaining: raise ValidationError('السداد المبكر والخصم يجب أن يساويا كامل المتبقي.')
            left=amount
            for p in parts:
                take=min(left,money(p['مبلغ القسط'])-p['paid_h']-p['waived_h'])
                if take:
                    p['paid_h']+=take; left-=take
                    allocated.append({'رقم التوزيع':f"{tr['رقم العملية']}:{p['رقم القسط']}:P",'رقم العملية':tr['رقم العملية'],'رقم القسط':p['رقم القسط'],'المبلغ':sar(take),'نوع التوزيع':'نقد','تاريخ العملية':d})
                if p['paid_h']+p['waived_h']==money(p['مبلغ القسط']) and p['paid_date'] is None: p['paid_date']=d
            if typ=='سداد مبكر':
                for p in parts:
                    gap=money(p['مبلغ القسط'])-p['paid_h']-p['waived_h']
                    if gap:
                        p['waived_h']+=gap; p['paid_date']=d
                        allocated.append({'رقم التوزيع':f"{tr['رقم العملية']}:{p['رقم القسط']}:W",'رقم العملية':tr['رقم العملية'],'رقم القسط':p['رقم القسط'],'المبلغ':sar(gap),'نوع التوزيع':'خصم','تاريخ العملية':d})
                settled.add(cid)
        elif money(tr['خصم التسوية']): raise ValidationError('الخصم مسموح في التسوية فقط.')
        running+=amount*CASH_TYPES[typ]
        if running<0: raise ValidationError('الحركة تجعل السيولة سالبة في تاريخها؛ راجع رأس المال وتاريخ الصرف.')
    ws=wb['توزيع التحصيل']; ws.delete_rows(2,ws.max_row)
    for r in allocated: append(wb,'توزيع التحصيل',r)
    # Formula cells are for Excel readers; the application calculates from raw records.
    cn=max(2,len(contracts)+1)
    for i,p in enumerate(schedule,2):
        ws=wb['الأقساط']
        for j,f in {
            6:f'=SUMIFS(\'توزيع التحصيل\'!$D$2:$D${max(2,len(allocated)+1)},\'توزيع التحصيل\'!$C$2:$C${max(2,len(allocated)+1)},A{i},\'توزيع التحصيل\'!$E$2:$E${max(2,len(allocated)+1)},"نقد")',
            7:f'=MAX(0,E{i}-F{i}-K{i})',8:f'=IF(COUNTIFS(\'العقود\'!$A$2:$A${cn},B{i},\'العقود\'!$J$2:$J${cn},"مصروف")=0,"لم يصرف",IF(G{i}=0,"مسدد",IF(D{i}<TODAY(),"متأخر",IF(D{i}=TODAY(),"مستحق اليوم",IF(F{i}>0,"سداد جزئي","قادم")))))' ,
            9:f'=IF(H{i}="لم يصرف",0,IF(G{i}>0,MAX(0,TODAY()-D{i}),IF(J{i}="",0,MAX(0,J{i}-D{i}))))',
            11:f'=SUMIFS(\'توزيع التحصيل\'!$D$2:$D${max(2,len(allocated)+1)},\'توزيع التحصيل\'!$C$2:$C${max(2,len(allocated)+1)},A{i},\'توزيع التحصيل\'!$E$2:$E${max(2,len(allocated)+1)},"خصم")',
        }.items(): ws.cell(i,j,f)
        p2=installments[p['رقم القسط']]; ws.cell(i,10).value=p2['paid_date']
    cn=max(2,len(contracts)+1); sn=max(2,len(schedule)+1); tn=max(2,len(rows(wb,'الحركات المالية'))+1)
    for i,c in enumerate(contracts.values(),2):
        parts=[r for r in installments.values() if r['رقم العقد']==c['رقم العقد']]
        rem=sum(money(r['مبلغ القسط'])-r['paid_h']-r['waived_h'] for r in parts)
        nextdue=min([day(r['تاريخ الاستحقاق']) for r in parts if money(r['مبلغ القسط'])>r['paid_h']+r['waived_h']],default=None)
        state='مسودة' if c['رقم العقد'] not in disbursed else ('مسدد مبكرًا' if c['رقم العقد'] in settled else ('مكتمل' if rem==0 else ('متأخر' if any(day(p['تاريخ الاستحقاق'])<today() and money(p['مبلغ القسط'])>p['paid_h']+p['waived_h'] for p in parts) else 'منتظم')))
        ws=wb['العقود']; ws.cell(i,10,'مصروف' if c['رقم العقد'] in disbursed else 'لم يصرف')
        for col,f in {
            11:f'=ROUNDDOWN(E{i}/F{i},2)',
            12:f'=SUMIFS(\'الأقساط\'!$F$2:$F${sn},\'الأقساط\'!$B$2:$B${sn},A{i})',
            13:f'=MAX(0,E{i}-L{i}-U{i})',14:f'=E{i}-D{i}',
            15:f'=COUNTIFS(\'الأقساط\'!$B$2:$B${sn},A{i},\'الأقساط\'!$G$2:$G${sn},0)',16:f'=F{i}-O{i}',
            17:f'=IF(E{i}=0,0,ROUND(D{i}*M{i}/E{i},2))',
            19:f'=EDATE(G{i},F{i}-1)',
            21:f'=SUMIFS(\'الحركات المالية\'!$J$2:$J${tn},\'الحركات المالية\'!$E$2:$E${tn},A{i},\'الحركات المالية\'!$K$2:$K${tn},"نشطة")',
        }.items(): ws.cell(i,col,f)
        ws.cell(i,18).value=nextdue
        ws.cell(i,20,f'=IF(J{i}="لم يصرف","مسودة",IF(M{i}=0,IF(COUNTIFS(\'الحركات المالية\'!$E$2:$E${tn},A{i},\'الحركات المالية\'!$D$2:$D${tn},"سداد مبكر",\'الحركات المالية\'!$K$2:$K${tn},"نشطة")>0,"مسدد مبكرًا","مكتمل"),IF(COUNTIFS(\'الأقساط\'!$B$2:$B${sn},A{i},\'الأقساط\'!$G$2:$G${sn},">0",\'الأقساط\'!$D$2:$D${sn},"<"&TODAY())>0,"متأخر","منتظم")))')
    # Use a computed snapshot for current internal delay counts and cash views.
    snap=_snapshot_data(wb,installments,disbursed,settled)
    for i,cl in enumerate(rows(wb,'العملاء'),2):
        ws=wb['العملاء']; cid=cl['رقم العميل']
        for col,source in [(15,'D'),(16,'M'),(17,'Q')]:
            ws.cell(i,col,f'=SUMIFS(\'العقود\'!${source}$2:${source}${cn},\'العقود\'!$B$2:$B${cn},A{i},\'العقود\'!$J$2:$J${cn},"مصروف")')
        ws.cell(i,18,sum(1 for x in snap['installments'] if x['client_id']==cid and x['disbursed'] and x['late_days']>0))
        ws.cell(i,19,f'=COUNTIFS(\'العقود\'!$B$2:$B${cn},A{i})')
    for i,tr in enumerate(rows(wb,'الحركات المالية'),2):
        wb['الحركات المالية'].cell(i,9,f'=IF(K{i}="نشطة",C{i}*{CASH_TYPES.get(tr["نوع العملية"],0)},0)')
    dash=wb['ملخص المحفظة']; dash.delete_rows(1,dash.max_row)
    dash.append(['إدارة التقسيط','']);dash.append(['المؤشر','القيمة بالريال / العدد'])
    formulas=[('السيولة المتاحة',f'=SUM(\'الحركات المالية\'!I2:I{tn})'),
      ('المتبقي للتحصيل',f'=SUMIFS(\'العقود\'!M2:M{cn},\'العقود\'!J2:J{cn},"مصروف")'),
      ('المتبقي من التمويل (تقدير نسبي)',f'=SUMIFS(\'العقود\'!Q2:Q{cn},\'العقود\'!J2:J{cn},"مصروف")'),
      ('إجمالي التمويل المصروف',f'=SUMIFS(\'الحركات المالية\'!C2:C{tn},\'الحركات المالية\'!D2:D{tn},"صرف تمويل",\'الحركات المالية\'!K2:K{tn},"نشطة")'),
      ('إجمالي التحصيل',f'=SUM(\'العقود\'!L2:L{cn})'),
      ('الربح التعاقدي المتوقع',f'=SUMIFS(\'العقود\'!N2:N{cn},\'العقود\'!J2:J{cn},"مصروف")'),
      ('خصومات التسوية',f'=SUM(\'العقود\'!U2:U{cn})'),
      ('العقود المصروفة',f'=COUNTIFS(\'العقود\'!J2:J{cn},"مصروف")'),
      ('عدد العملاء',f'=COUNTA(\'العملاء\'!A2:A{max(2,len(clients)+1)})')]
    for r in formulas: dash.append(r)
    dash.append(['المصدر',SOURCE]);dash.append(['ملاحظة','الربح المتوقع لا يعني الربح المحقق؛ التصنيف يدوي.'])
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    return snap

def _snapshot_data(wb,inst,disbursed,settled,as_of=None):
    now=as_of or today(); cs={r['رقم العقد']:r for r in rows(wb,'العقود')}; ins=[]; cons=[]
    for p in inst.values():
        total=money(p['مبلغ القسط']); rem=total-p['paid_h']-p['waived_h']; due=day(p['تاريخ الاستحقاق']); funded=p['رقم العقد'] in disbursed
        late=max(0,((p['paid_date'] or now)-due).days) if funded else 0
        status='لم يصرف' if not funded else ('مسدد' if rem==0 else ('متأخر' if due<now else ('مستحق اليوم' if due==now else ('سداد جزئي' if p['paid_h'] else 'قادم'))))
        ins.append({**p,'remaining_h':rem,'total_h':total,'due':due,'late_days':late,'status':status,'disbursed':funded,'client_id':cs[p['رقم العقد']]['رقم العميل']})
    for cid,c in cs.items():
        pp=[x for x in ins if x['رقم العقد']==cid]; rem=sum(x['remaining_h'] for x in pp); paid=sum(x['paid_h'] for x in pp); funded=cid in disbursed
        state='مسودة' if not funded else ('مسدد مبكرًا' if cid in settled else ('مكتمل' if rem==0 else ('متأخر' if any(x['status']=='متأخر' for x in pp) else 'منتظم')))
        p=money(c['مبلغ التمويل']); total=money(c['إجمالي السداد']); discount=sum(x['waived_h'] for x in pp)
        cons.append({**c,'client_id':c['رقم العميل'],'principal_h':p,'total_h':total,'paid_h':paid,'remaining_h':rem,'discount_h':discount,
           'principal_remaining_h':money(Decimal(p)*Decimal(rem)/Decimal(total)/100) if total else 0,
           'profit_collected_h':paid-money(Decimal(p)*Decimal(paid)/Decimal(total)/100) if total else 0,
           'disbursed':funded,'status':state,'next_due':min([x['due'] for x in pp if x['remaining_h']],default=None),
           'overdue_h':sum(x['remaining_h'] for x in pp if x['due']<now) if funded else 0,'paid_count':sum(x['remaining_h']==0 for x in pp)})
    return {'contracts':cons,'installments':ins,'cash_h':cash(wb),'clients':rows(wb,'العملاء'),'transactions':rows(wb,'الحركات المالية'),'followups':rows(wb,'متابعة العملاء')}

def snapshot(wb,as_of=None):
    # Reconstruct from authoritative ledger, independent of Excel formula caches.
    if as_of is None: return rebuild(wb)
    snap=rebuild(wb)
    inst={x['رقم القسط']:x for x in snap['installments']}
    dis={x['رقم العقد']:True for x in snap['contracts'] if x['disbursed']}
    settled={x['رقم العقد'] for x in snap['contracts'] if x['status']=='مسدد مبكرًا'}
    return _snapshot_data(wb,inst,dis,settled,as_of)

def validate_book(wb):
    for name,h in SCHEMA.items():
        if name not in wb.sheetnames or [c.value for c in wb[name][1]][:len(h)]!=h: raise ValidationError(f'بنية ورقة {name} غير مطابقة.')
    if str(settings(wb).get('schema_version'))!='1': raise ValidationError('نسخة بنية الملف غير مدعومة.')
    for sheet,key in [('العملاء','رقم العميل'),('العقود','رقم العقد'),('متابعة العملاء','رقم المتابعة')]:
        rr=rows(wb,sheet)
        if len({r[key] for r in rr})!=len(rr): raise ValidationError('معرفات سجلات مكررة.')
    identities=[text(r['رقم الهوية']) for r in rows(wb,'العملاء') if r['رقم الهوية']]
    if len(identities)!=len(set(identities)): raise ValidationError('أرقام هويات مكررة.')
    for c in rows(wb,'العملاء'):
        if not text(c['اسم العميل']) or not text(c['رقم الجوال']): raise ValidationError('اسم وجوال العميل مطلوبان.')
        if c['حالة العميل'] not in CLIENT_STATES or c['التصنيف الائتماني'] not in GRADES: raise ValidationError('حالة أو تصنيف عميل غير صالح.')
        if money(c['الحد الائتماني'])<0 or money(c['الدخل الشهري'])<0: raise ValidationError('بيانات العميل المالية لا تقبل السالب.')
    for c in rows(wb,'العقود'):
        n=c['مدة العقد']
        if not isinstance(n,(int,float)) or int(n)!=n or not 1<=n<=120: raise ValidationError('مدة عقد غير صالحة.')
        if money(c['مبلغ التمويل'])<=0 or money(c['إجمالي السداد'])<money(c['مبلغ التمويل']) or money(c['رسوم العقد'])<0:
            raise ValidationError('مبالغ عقد غير صالحة.')
        if day(c['تاريخ العقد']) is None or day(c['أول قسط']) is None or day(c['أول قسط'])<day(c['تاريخ العقد']):
            raise ValidationError('تواريخ عقد غير صالحة.')
    contracts={c['رقم العقد'] for c in rows(wb,'العقود')}
    clients={c['رقم العميل'] for c in rows(wb,'العملاء')}
    for p in rows(wb,'الأقساط'):
        if p['رقم العقد'] not in contracts or day(p['تاريخ الاستحقاق']) is None or money(p['مبلغ القسط'])<=0:
            raise ValidationError('قسط غير صالح أو مرتبط بعقد مفقود.')
    for f in rows(wb,'متابعة العملاء'):
        if f['رقم العميل'] not in clients or (f['رقم العقد'] and f['رقم العقد'] not in contracts):
            raise ValidationError('متابعة مرتبطة بسجل مفقود.')
    requests=[r['رقم الطلب'] for r in rows(wb,'الحركات المالية') if r['رقم الطلب']]
    if len(requests)!=len(set(requests)): raise ValidationError('أرقام طلبات مالية مكررة.')
    if any(r['حالة العملية'] not in ['نشطة','ملغاة'] for r in rows(wb,'الحركات المالية')):
        raise ValidationError('حالة حركة مالية غير صالحة.')
    rebuild(wb)

def format_book(wb):
    for ws in wb:
        ws.sheet_view.rightToLeft=True; ws.sheet_view.showGridLines=False
        ws.freeze_panes='C2' if ws.title!='ملخص المحفظة' else None
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.sheet_properties.outlinePr.summaryRight=False
        ws.print_options.horizontalCentered=True
        ws.sheet_properties.tabColor='178D88' if ws.title=='ملخص المحفظة' else '344B55'
        for r in ws:
            for c in r:
                c.font=Font(name='Arial',size=11,color='112F39')
                c.alignment=Alignment(horizontal='right',vertical='center',readingOrder=2)
                if c.data_type=='f': c.font=Font(name='Arial',size=11,color='13795B')
                elif isinstance(c.value,(float,int)): c.font=Font(name='Arial',size=11,color='2358A5');c.number_format='#,##0.00;(#,##0.00);"—"'
                elif isinstance(c.value,(date,datetime)): c.number_format='yyyy-mm-dd'
        for c in ws[1]: c.fill=PatternFill('solid',fgColor='20343E');c.font=Font(name='Arial',bold=True,color='FFFFFF',size=11)
        ws.row_dimensions[1].height=30
        for j in range(1,ws.max_column+1):
            header=ws.cell(1,j).value or ''; width=24
            if any(k in str(header) for k in ['ملاحظات','تفاصيل','الوصف']): width=48
            if header in ['رقم الهوية','رقم الجوال','جوال الكفيل']: width=21
            ws.column_dimensions[get_column_letter(j)].width=width
        if ws.title in SCHEMA and ws.max_row>1:
            for name in list(ws.tables): del ws.tables[name]
            t=Table(displayName='Table'+str(wb.sheetnames.index(ws.title)),ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}')
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(t)
        ws.auto_filter.ref=ws.dimensions if ws.max_row>1 else None
    d=wb['ملخص المحفظة'];d.column_dimensions['A'].width=43;d.column_dimensions['B'].width=108
    d['A1'].font=Font(name='Arial',size=17,bold=True,color='153945');d['A1'].fill=PatternFill(fill_type=None)
    d['B1'].fill=PatternFill(fill_type=None)
    for c in d[2]: c.fill=PatternFill('solid',fgColor='E2F0F0');c.font=Font(name='Arial',bold=True,color='153945',size=11)
    for row in range(3,12): d.row_dimensions[row].height=27;d.cell(row,2).number_format='#,##0.00;(#,##0.00);"—"'
    d['B10'].number_format='0';d['B11'].number_format='0'
    for col in ['C','D','L']:
        wb['العملاء'].column_dimensions[col].width=24
    config=wb['الإعدادات']
    config.column_dimensions['B'].width=90
    config.column_dimensions['C'].width=68
    for r in range(2,config.max_row+1):
        config.cell(r,2).alignment=Alignment(horizontal='right',vertical='center',wrap_text=True,readingOrder=2)
        config.cell(r,3).alignment=Alignment(horizontal='right',vertical='center',wrap_text=True,readingOrder=2)
        config.row_dimensions[r].height=45
    config['B2'].number_format='0'
    config['B3'].number_format='0'
    config['B4'].number_format='0.0%'

def dump(wb):
    b=BytesIO();wb.save(b);return b.getvalue()
def load(data):
    wb=load_workbook(BytesIO(data));validate_book(wb);return wb

def perform(wb,action,data,request_id=None):
    """Mutate a disposable copy; callers publish it atomically only after validation."""
    request_id=request_id or uid('REQ')
    if any(r['رقم السجل']==request_id for r in rows(wb,'سجل التغييرات')): return 'سبق تنفيذ الطلب'
    if action=='add_client': ident=add_client(wb,data)
    elif action=='edit_client': ident=data['id'];edit_client(wb,ident,data)
    elif action=='add_contract': ident=add_contract(wb,data)
    elif action=='edit_contract': ident=data['id'];edit_contract(wb,ident,data)
    elif action=='transaction': ident=tx(wb,request_id=request_id,**data)
    elif action=='edit_transaction':
        old=require(wb,'الحركات المالية','رقم العملية',data['id'])
        if old['حالة العملية']!='نشطة': raise ValidationError('الحركة غير نشطة.')
        update(wb,'الحركات المالية','رقم العملية',data['id'],{'حالة العملية':'ملغاة','إلغاء بواسطة':request_id})
        ident=tx(wb,request_id=request_id,**data['replacement'])
    elif action in ['void_transaction','restore_transaction']:
        ident=data['id'];old=require(wb,'الحركات المالية','رقم العملية',ident)
        target='ملغاة' if action=='void_transaction' else 'نشطة'
        if old['حالة العملية']==target: raise ValidationError('الحركة بهذه الحالة بالفعل.')
        update(wb,'الحركات المالية','رقم العملية',ident,{'حالة العملية':target,'إلغاء بواسطة':request_id if target=='ملغاة' else ''})
    elif action=='add_followup': ident=add_followup(wb,data)
    elif action=='edit_followup':
        ident=data['id'];old=require(wb,'متابعة العملاء','رقم المتابعة',ident)
        changes={k:text(data.get(k,old.get(k))) for k in ['نتيجة التواصل','ملاحظات','الحالة']}
        if changes['الحالة'] not in ['مفتوحة','مكتملة']: raise ValidationError('حالة متابعة غير صالحة.')
        for k in ['وعد السداد','المتابعة القادمة']: changes[k]=day(data.get(k,old.get(k)))
        amount=money(data.get('المبلغ الموعود',old.get('المبلغ الموعود')))
        if amount<0: raise ValidationError('مبلغ الوعد لا يقبل السالب.')
        changes['المبلغ الموعود']=sar(amount);changes['آخر تعديل']=stamp()
        update(wb,'متابعة العملاء','رقم المتابعة',ident,changes)
    elif action=='close_followup':
        ident=data['id'];update(wb,'متابعة العملاء','رقم المتابعة',ident,{'الحالة':'مكتملة','آخر تعديل':stamp()})
    elif action=='settings':
        n=int(data['default_months']);rate=float(data['default_profit_rate'])
        if not 1<=n<=120 or not 0<=rate<=10: raise ValidationError('الإعدادات خارج النطاق.')
        for k,v in [('default_months',n),('default_profit_rate',rate)]: update(wb,'الإعدادات','المفتاح',k,{'القيمة':v})
        ident='settings'
    else: raise ValidationError('الإجراء غير مدعوم.')
    validate_book(wb)
    append(wb,'سجل التغييرات',{'رقم السجل':request_id,'الوقت':stamp(),'نوع الإجراء':action,'رقم السجل المتأثر':ident,'تفاصيل':'تم التحقق وإعادة احتساب الأرصدة'})
    format_book(wb)
    return ident
