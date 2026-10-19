"""Installment System 2.0. One engine and UI entrypoint; XLSX is authoritative.

Money is calculated in integer halalas. Transactions are replayed, not trusted
from cached formula values. Executed terms are immutable; corrections void a
whole posting group and replay every subsequent group before publication.
"""
from __future__ import annotations
from calendar import monthrange
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zoneinfo import ZoneInfo
import hashlib
import json
import re
import os
import tempfile
from urllib.parse import quote

import pandas as pd
import requests
import streamlit as st
from filelock import FileLock
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

VERSION = '2.3.0'
FORMAT = 'IBRAHIM_INSTALLMENTS_V2'
ROOT = Path(__file__).resolve().parent
TZ = ZoneInfo('Asia/Riyadh')
SELF, EXTERNAL, MANAGED = 'تمويل ذاتي', 'استثمار عبر مكتب', 'إدارة تمويل مستثمر'
MODES = [SELF, EXTERNAL, MANAGED]
HEADERS = {
 'العملاء': ['المعرف','الاسم','الجوال','الهوية','المستوى','الحالة','الحد الائتماني','العنوان','جهة العمل','الدخل','الكفيل','جوال الكفيل','ملاحظات'],
 'المستثمرون': ['المعرف','الاسم','الجوال','الهوية','الحالة','ملاحظات'],
 'المكاتب': ['المعرف','الاسم','الجوال','الحالة','ملاحظات'],
 'العقود': ['المعرف','العميل','النوع','المستثمر','المكتب','تاريخ العقد','أول استحقاق','الأشهر','تكلفة الشراء','سعر الكاش','إجمالي التقسيط','التمويل المحول','رسوم مكتب مدفوعة','رسوم مكتب ثابتة','رسوم العميل','مصاريف أخرى','ملاحظات','فرق الشراء','تكلفة المستثمر','ربح المستثمر المتوقع','ربح المكتب المتوقع','عائد المستثمر','المحصل','الخصم','المتبقي','الحالة','فائض لدى المكتب','تاريخ التنفيذ'],
 'الأقساط': ['المعرف','العقد','التسلسل','الاستحقاق','المبلغ','المحصل','الخصم','المتبقي','الحالة','أيام التأخير','تاريخ السداد'],
 'الحركات': ['المعرف','المجموعة','التاريخ','الوقت','النوع','العقد','المستثمر','المبلغ','أثر البنك','أثر أموال المستثمر','خصم التسوية','المرجع','ملاحظات','الحالة','رقم الطلب','الرصيد البنكي','أموال المستثمرين','السيولة الحرة'],
 'المتابعات': ['المعرف','العميل','العقد','تاريخ التواصل','الوسيلة','النتيجة','وعد السداد','المبلغ الموعود','المتابعة القادمة','الحالة','ملاحظات'],
 'النظام': ['المفتاح','القيمة'],
}
HEADERS['تفاصيل العقود']=['المعرف','نوع السلعة','وصف السلعة','الدفعة الأولى','قيمة سند لأمر','ملغى','سبب الإلغاء','تاريخ الإلغاء']
INPUT_COLUMNS = {'العقود':17,'الحركات':15}
DEFAULTS = {'format':FORMAT,'version':VERSION,'default_months':16,'default_profit_rate':0.60}

class Error(ValueError): pass
class Conflict(Error): pass
def today(): return datetime.now(TZ).date()
def uid(prefix): return prefix+'-'+uuid4().hex[:12].upper()
def money(v):
    try:
        x=Decimal(str(v if v not in [None,''] else 0))
        if not x.is_finite(): raise ValueError()
        return int((x*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
    except Exception: raise Error('أدخل مبلغًا رقميًا صالحًا.')
def sar(h): return float(Decimal(h)/100)
def dt(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    if not v: return None
    try: return date.fromisoformat(str(v)[:10])
    except ValueError: raise Error('التاريخ غير صالح.')
def month(d,n):
    y,m=divmod(d.year*12+d.month-1+n,12);m+=1
    return date(y,m,min(d.day,monthrange(y,m)[1]))
def txt(v):
    v=str(v or '').strip()
    if len(v)>4000 or any(ord(c)<32 and c not in '\n\t\r' for c in v):raise Error('النص غير صالح أو طويل جدًا.')
    return v
def empty(): return {**{k:[] for k in HEADERS},'النظام':[{'المفتاح':k,'القيمة':v} for k,v in DEFAULTS.items()]}
def lookup(data,sheet,ident):
    r=next((x for x in data[sheet] if x['المعرف']==ident),None)
    if r is None:raise Error('السجل المرتبط غير موجود: '+str(ident))
    return r
def setting(data,k): return next((r['القيمة'] for r in data['النظام'] if r['المفتاح']==k),None)

def pricing(c):
    a={k:money(c.get(k)) for k in HEADERS['العقود'][8:16]}
    purchase=a['تكلفة الشراء'];cash=a['سعر الكاش'];total=a['إجمالي التقسيط'];fund=a['التمويل المحول']
    managed=c['النوع']==MANAGED
    cost=purchase+(a['رسوم مكتب ثابتة'] if managed else a['رسوم مكتب مدفوعة']+a['مصاريف أخرى'])
    investor_profit=total-cost
    office_profit=(a['رسوم مكتب ثابتة']+a['رسوم العميل']-a['مصاريف أخرى']) if managed else investor_profit+a['رسوم العميل']
    return {'purchase':purchase,'cash':cash,'total':total,'fund':fund,'paid_fee':a['رسوم مكتب مدفوعة'],
      'fee':a['رسوم مكتب ثابتة'],'client_fee':a['رسوم العميل'],'expenses':a['مصاريف أخرى'],
      'saving':cash-purchase,'surplus':fund-purchase if c['النوع']!=SELF else 0,
      'cost':cost,'investor_profit':investor_profit,'office_profit':office_profit,
      'roi':float(Decimal(investor_profit)/cost) if cost else 0}

def details(data,cid):
    return next((r for r in data.get('تفاصيل العقود',[]) if r['المعرف']==cid),{})

def digits(value):
    return str(value or '').strip().translate(str.maketrans('٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹','01234567890123456789'))

def valid_person(p,require_id=False):
    phone=digits(p.get('الجوال'))
    if not re.fullmatch(r'05[0-9]{8}',phone):raise Error('رقم الجوال يجب أن يكون 10 أرقام ويبدأ بـ05، مثال 0512345678.')
    identity=digits(p.get('الهوية'))
    if (require_id or identity) and not re.fullmatch(r'[0-9]{10}',identity):raise Error('رقم الهوية يجب أن يكون 10 أرقام بالضبط.')
    guarantor=digits(p.get('جوال الكفيل'))
    if guarantor and not re.fullmatch(r'05[0-9]{8}',guarantor):raise Error('جوال الكفيل يجب أن يكون 10 أرقام ويبدأ بـ05.')
    return {**p,'الجوال':phone,**({'الهوية':identity} if 'الهوية' in p or require_id else {}),**({'جوال الكفيل':guarantor} if 'جوال الكفيل' in p else {})}

def investor_capital(data,snap,ident):
    contracts=[s for s in snap['contracts'] if s['active'] and s['c'].get('المستثمر')==ident]
    return [('إجمالي رأس المال الممول',sum(s['p']['fund'] for s in contracts)),
            ('تكلفة السلع المشتراة',sum(s['p']['purchase'] for s in contracts)),
            ('أمانات نقدية متاحة',snap['holdings'].get(ident,0)),
            ('أقساط متبقية للمستثمر',sum(s['remaining'] for s in contracts))]

def valid_contract(data,c):
    client=lookup(data,'العملاء',c['العميل'])
    if c['النوع'] not in MODES:raise Error('اختر نوع الصفقة.')
    p=pricing(c)
    if min(p['purchase'],p['cash'],p['fund'])<=0 or p['total']<p['cash']:raise Error('تحقق من تكلفة الشراء وسعر الكاش وإجمالي التقسيط.')
    if min(p['paid_fee'],p['fee'],p['client_fee'],p['expenses'])<0:raise Error('الرسوم والمصاريف لا تقبل السالب.')
    if c['النوع']==SELF and p['fund']!=p['purchase']:raise Error('التمويل الذاتي يساوي تكلفة الشراء الفعلية.')
    if c['النوع']!=SELF and p['fund']<p['purchase']:raise Error('المبلغ المحول أقل من تكلفة الشراء.')
    if c['النوع']==MANAGED:
        lookup(data,'المستثمرون',c['المستثمر'])
        if p['paid_fee']:raise Error('عقد المستثمر يستخدم رسوم المكتب الثابتة.')
    elif c.get('المستثمر') or p['fee']:raise Error('المستثمر والرسوم الثابتة يخصان العقود المدارة فقط.')
    if c['النوع']==EXTERNAL:lookup(data,'المكاتب',c['المكتب'])
    if c['النوع']!=EXTERNAL and c.get('المكتب'):raise Error('المكتب الخارجي يخص الاستثمار عبر مكتب.')
    n=c['الأشهر']
    if not isinstance(n,(int,float)) or int(n)!=n or not 1<=n<=120 or p['total']<n:raise Error('مدة العقد من شهر إلى 120 شهرًا.')
    if not dt(c['تاريخ العقد']) or dt(c['تاريخ العقد'])>today() or not dt(c['أول استحقاق']) or dt(c['أول استحقاق'])<dt(c['تاريخ العقد']):raise Error('تواريخ العقد غير صالحة.')
    return client,p

def snapshot(data,as_of=None):
    """Replay all groups in date order. Void is all-or-nothing for each group."""
    now=as_of or today();ids={};contracts={};parts=[];states={};ledger=[]
    for sheet in ['العملاء','المستثمرون','المكاتب','العقود','الحركات','المتابعات']:
        vals=[r.get('المعرف') for r in data[sheet]]
        if any(not x for x in vals) or len(set(vals))!=len(vals):raise Error('معرف مفقود أو مكرر في '+sheet)
    for sheet in ['العملاء','المستثمرون','المكاتب']:
        for r in data[sheet]:
            if not txt(r.get('الاسم')) or not txt(r.get('الجوال')) or r.get('الحالة') not in ['نشط','موقوف','مؤرشف']:raise Error('بيانات غير صالحة في '+sheet)
        identities=[txt(r.get('الهوية')) for r in data[sheet] if txt(r.get('الهوية'))]
        if len(set(identities))!=len(identities):raise Error('الهوية مكررة في '+sheet)
    for r in data['العملاء']:
        if r.get('المستوى') not in ['غير مصنف','A','B','C','D'] or min(money(r.get('الحد الائتماني')),money(r.get('الدخل')))<0:raise Error('تصنيف أو بيانات مالية للعميل غير صالحة.')
    metas=data.get('تفاصيل العقود',[])
    mids=[r['المعرف'] for r in metas]
    if len(mids)!=len(set(mids)) or any(x not in {c['المعرف'] for c in data['العقود']} for x in mids):raise Error('تفاصيل عقد مكررة أو مرتبطة بعقد مفقود.')
    for c in data['العقود']:
        _,p=valid_contract(data,c);cid=c['المعرف'];contracts[cid]=c
        meta=details(data,cid);down=money(meta.get('الدفعة الأولى'));note=money(meta.get('قيمة سند لأمر'))
        if down<0 or down>=p['total'] or note<0:raise Error('الدفعة الأولى من صفر إلى أقل من سعر البيع، وقيمة السند غير سالبة.')
        n=int(c['الأشهر']);net=p['total']-down
        if net<n:raise Error('المبلغ بعد الدفعة الأولى أقل من عدد الأقساط بالهللات.')
        base=net//n
        schedule=[{'id':cid+'-'+str(i+1).zfill(3),'contract':cid,'sequence':i+1,'due':month(dt(c['أول استحقاق']),i),
          'amount':base if i<n-1 else net-base*(n-1),'paid':0,'waived':0,'paid_date':None} for i in range(n)]
        states[cid]={'c':c,'p':p,'parts':schedule,'active':False,'activated':None,'settled':False,'refund':0,'down':down,'cancelled':bool(meta.get('ملغى'))}
        parts.extend(schedule)
    bank=0;holdings={r['المعرف']:0 for r in data['المستثمرون']};refs=set();requests_seen={};groups={}
    for tr in data['الحركات']:
        if tr.get('الحالة') not in ['نشطة','ملغاة']:raise Error('حالة حركة غير صالحة.')
        if not dt(tr['التاريخ']) or dt(tr['التاريخ'])>today():raise Error('تاريخ الحركة غير صالح.')
        group=txt(tr.get('المجموعة'));req=txt(tr.get('رقم الطلب'))
        if not group or not req:raise Error('مجموعة الحركة ورقم الطلب مطلوبان.')
        if req in requests_seen and requests_seen[req]!=group:raise Error('رقم طلب مكرر.')
        requests_seen[req]=group
        groups.setdefault(group,[]).append(tr)
    ordered=sorted(groups.values(),key=lambda ts:(dt(ts[0]['التاريخ']),str(ts[0]['الوقت']),ts[0]['المجموعة']))
    for ts in ordered:
        if len({(str(t['التاريخ']),t['الوقت'],t['الحالة'],t['رقم الطلب'],t.get('العقد') or '') for t in ts})!=1:raise Error('مجموعة حركة غير مكتملة.')
        if ts[0]['الحالة']=='ملغاة':continue
        kinds=[t['النوع'] for t in ts];cid=ts[0].get('العقد') or '';s=states.get(cid);c=contracts.get(cid)
        if cid and not s:raise Error('حركة مرتبطة بعقد مفقود.')
        if c and dt(ts[0]['التاريخ'])<dt(c['تاريخ العقد']):raise Error('الحركة تسبق العقد.')
        if 'تنفيذ' in kinds:
            if s['active']:raise Error('تنفيذ مكرر للعقد.')
            if s['cancelled']:raise Error('العقد ملغى؛ لا يمكن تنفيذه.')
            expected=activation_rows(c,details(data,cid))
            actual=[(t['النوع'],money(t['المبلغ']),money(t['أثر البنك']),money(t['أثر أموال المستثمر']),t.get('المستثمر') or '') for t in ts]
            exp=[(t[0],t[1],t[2],t[3],c.get('المستثمر') or '') for t in expected]
            if actual!=exp:raise Error('قيود التنفيذ لا تطابق شروط العقد.')
            client=lookup(data,'العملاء',c['العميل']);limit=money(client.get('الحد الائتماني'))
            outstanding=sum(x['p']['cash']*sum(z['amount']-z['paid']-z['waived'] for z in x['parts'])//x['p']['total'] for x in states.values() if x['active'] and x['c']['العميل']==c['العميل'])
            if limit and outstanding+s['p']['cash']>limit:raise Error('التنفيذ يتجاوز الحد الائتماني للعميل.')
            s['active']=True;s['activated']=dt(ts[0]['التاريخ'])
        elif len(ts)!=1:raise Error('مجموعة مالية غير معروفة.')
        for tr in ts:
            kind=tr['النوع'];amount=money(tr['المبلغ']);bd=money(tr['أثر البنك']);hd=money(tr['أثر أموال المستثمر']);inv=tr.get('المستثمر') or ''
            if amount<0:raise Error('مبلغ الحركة لا يقبل السالب.')
            if kind in ['سداد','تسوية']:
                if not s['active'] or s['settled'] or amount<=0:raise Error('السداد يتطلب عقدًا منفذًا وغير مسوى.')
                rem=sum(x['amount']-x['paid']-x['waived'] for x in s['parts']);waive=money(tr.get('خصم التسوية'))
                if waive<0 or amount>rem or (kind=='سداد' and waive) or (kind=='تسوية' and amount+waive!=rem):raise Error('مبلغ السداد أو الخصم لا يطابق المتبقي.')
                left=amount
                for z in s['parts']:
                    take=min(left,z['amount']-z['paid']-z['waived']);z['paid']+=take;left-=take
                    if z['paid']==z['amount'] and z['paid_date'] is None:z['paid_date']=dt(tr['التاريخ'])
                if kind=='تسوية':
                    for z in s['parts']:
                        z['waived']=z['amount']-z['paid']
                        if z['paid_date'] is None:z['paid_date']=dt(tr['التاريخ'])
                    s['settled']=True
                expected_hd=amount if c['النوع']==MANAGED else 0
                if bd!=amount or hd!=expected_hd or inv!=(c.get('المستثمر') or ''):raise Error('ملكية السداد غير مطابقة للعقد.')
            elif kind=='دفعة أولى':
                if amount!=s['down'] or bd!=amount or hd!=(amount if c['النوع']==MANAGED else 0) or inv!=(c.get('المستثمر') or ''):raise Error('قيد الدفعة الأولى غير مطابق.')
            elif kind=='استرداد فائض':
                if not s or not s['active'] or c['النوع']!=EXTERNAL or amount<=0:raise Error('استرداد الفائض يتطلب استثمارًا منفذًا عبر مكتب.')
                s['refund']+=amount
                if s['refund']>s['p']['surplus'] or bd!=amount or hd or inv:raise Error('استرداد الفائض يتجاوز المستحق أو قيده غير مطابق.')
            elif kind=='تحويل للمستثمر':
                if cid or inv not in holdings or bd!=-amount or hd!=-amount or amount<=0:raise Error('تحويل المستثمر غير صالح.')
            elif kind in ['إضافة رأس مال','سحب رأس مال','مصروف مكتب']:
                if cid or inv or hd or bd!=(amount if kind=='إضافة رأس مال' else -amount) or amount<=0:raise Error('قيد السيولة غير مطابق.')
            elif 'تنفيذ' not in kinds:raise Error('نوع حركة غير معروف.')
            if kind!='تسوية' and money(tr.get('خصم التسوية')):raise Error('الخصم يخص التسوية فقط.')
            ref=txt(tr.get('المرجع'))
            if ref:
                if ref in refs:raise Error('مرجع مالي مكرر.')
                refs.add(ref)
            bank+=bd
            if inv:
                if inv not in holdings:raise Error('مستثمر الحركة غير موجود.')
                holdings[inv]+=hd
            elif hd:raise Error('أثر المستثمر دون معرف مستثمر.')
            if bank<0 or any(v<0 for v in holdings.values()) or bank-sum(holdings.values())<0:raise Error('رصيد غير كافٍ أو العملية تستخدم أموال المستثمر لصالح المكتب.')
            ledger.append({**tr,'bank':bank,'held':sum(holdings.values()),'free':bank-sum(holdings.values())})
    summaries=[]
    for cid,s in states.items():
        pp=s['parts'];p=s['p'];paid=sum(z['paid'] for z in pp)+(s['down'] if s['active'] else 0);waived=sum(z['waived'] for z in pp);rem=p['total']-paid-waived
        overdue=sum(z['amount']-z['paid']-z['waived'] for z in pp if z['due']<now) if s['active'] else 0
        if s['cancelled'] and s['active']:raise Error('العقد الملغى لا يقبل حركات مالية نشطة.')
        if s['cancelled']:rem=0
        state='ملغى' if s['cancelled'] else 'مسودة' if not s['active'] else ('مسدد مبكرًا' if s['settled'] else ('مكتمل' if rem==0 else ('متأخر' if overdue else 'منتظم')))
        summaries.append({**s,'id':cid,'paid':paid,'waived':waived,'remaining':rem,'overdue':overdue,'status':state,
          'surplus_due':p['surplus']-s['refund'] if s['active'] and s['c']['النوع']==EXTERNAL else 0,
          'next_due':min((z['due'] for z in pp if z['amount']>z['paid']+z['waived']),default=None) if s['active'] else None})
        for z in pp:
            z['remaining']=0 if s['cancelled'] else z['amount']-z['paid']-z['waived'];z['active']=s['active'];z['client']=s['c']['العميل']
            z['late_days']=max(0,((z['paid_date'] or now)-z['due']).days) if s['active'] else 0
            z['status']='ملغى' if s['cancelled'] else 'لم ينفذ' if not s['active'] else ('مسدد' if z['remaining']==0 else ('متأخر' if z['due']<now else ('مستحق اليوم' if z['due']==now else ('سداد جزئي' if z['paid'] else 'قادم'))))
    for r in data['المتابعات']:
        lookup(data,'العملاء',r['العميل'])
        if r.get('العقد') and lookup(data,'العقود',r['العقد'])['العميل']!=r['العميل']:raise Error('متابعة مرتبطة بعميل مختلف.')
        if r.get('الحالة') not in ['مفتوحة','مكتملة'] or money(r.get('المبلغ الموعود'))<0:raise Error('متابعة غير صالحة.')
    return {'contracts':summaries,'parts':parts,'ledger':ledger,'bank':bank,'holdings':holdings,'held':sum(holdings.values()),'free':bank-sum(holdings.values())}

def activation_rows(c,meta=None):
    p=pricing(c);mode=c['النوع']
    # Investor pays the fixed fee in addition to funding. Transfer to office is
    # an internal ownership transfer: bank 0, investor-held cash -fee.
    if mode==MANAGED:
        a=[('تنفيذ',p['fund']+p['fee'],p['fund']+p['fee'],p['fund']+p['fee']),
           ('شراء للمستثمر',p['purchase'],-p['purchase'],-p['purchase'])]
        if p['fee']:a.append(('رسوم ثابتة',p['fee'],0,-p['fee']))
    else:
        out=(p['fund'] if mode==EXTERNAL else p['purchase'])+p['paid_fee']+p['expenses']
        a=[('تنفيذ',out,-out,0)]
    if p['client_fee']:a.append(('رسوم عميل',p['client_fee'],p['client_fee'],0))
    if mode==MANAGED and p['expenses']:a.append(('مصروف صفقة',p['expenses'],-p['expenses'],0))
    down=money((meta or {}).get('الدفعة الأولى'))
    if down:a.append(('دفعة أولى',down,down,down if mode==MANAGED else 0))
    return a

def post(data,rows,when,cid='',investor='',ref='',notes='',request=None,discount=0):
    when=dt(when)
    if not when or when>today():raise Error('اختر تاريخًا لا يتجاوز اليوم.')
    group=uid('G');request=request or uid('REQ');stamp=datetime.now(TZ).isoformat(timespec='microseconds')
    for i,(kind,amount,bd,hd) in enumerate(rows):
        data['الحركات'].append({'المعرف':uid('TX'),'المجموعة':group,'التاريخ':when,'الوقت':stamp,'النوع':kind,'العقد':cid,'المستثمر':investor,
          'المبلغ':sar(amount),'أثر البنك':sar(bd),'أثر أموال المستثمر':sar(hd),'خصم التسوية':sar(discount if i==0 else 0),
          'المرجع':txt(ref) if i==0 else '', 'ملاحظات':txt(notes),'الحالة':'نشطة','رقم الطلب':request})
    return group

def perform(original,action,payload,request=None):
    data=deepcopy(original);data.setdefault('تفاصيل العقود',[]);p=deepcopy(payload);request=request or uid('REQ')
    if any(t['رقم الطلب']==request for t in data['الحركات']):return data
    if action in ['save_client','save_investor','save_office']:
        p=valid_person(p,action=='save_client')
        sheet={'save_client':'العملاء','save_investor':'المستثمرون','save_office':'المكاتب'}[action]
        ident=p.get('المعرف') or uid({'العملاء':'CL','المستثمرون':'INV','المكاتب':'OF'}[sheet]);p['المعرف']=ident
        old=next((r for r in data[sheet] if r['المعرف']==ident),None)
        if p.get('الحالة')=='مؤرشف':
            key={'العملاء':'العميل','المستثمرون':'المستثمر','المكاتب':'المكتب'}[sheet]
            if any(s['c'].get(key)==ident and s['active'] and s['remaining']>0 for s in snapshot(data)['contracts']):raise Error('السجل لديه عقد قائم. استخدم موقوف بدل الأرشفة.')
        if old:old.update(p)
        else:data[sheet].append(p)
    elif action=='save_contract':
        extra=p.pop('details',None)
        if extra is not None and not txt(extra.get('نوع السلعة')):raise Error('نوع السلعة مطلوب.')
        ident=p.get('المعرف') or uid('CT');p['المعرف']=ident
        if details(data,ident).get('ملغى'):raise Error('العقد ملغى.')
        if any(t.get('العقد')==ident and t['الحالة']=='نشطة' for t in data['الحركات']):raise Error('شروط العقد المنفذ مقفلة. ألغِ التنفيذ والحركات التابعة أولًا.')
        if lookup(data,'العملاء',p['العميل'])['الحالة']!='نشط':raise Error('العميل غير نشط.')
        if p['النوع']==MANAGED and lookup(data,'المستثمرون',p['المستثمر'])['الحالة']!='نشط':raise Error('المستثمر غير نشط.')
        if p['النوع']==EXTERNAL and lookup(data,'المكاتب',p['المكتب'])['الحالة']!='نشط':raise Error('المكتب غير نشط.')
        old=next((r for r in data['العقود'] if r['المعرف']==ident),None)
        if old:old.update(p)
        else:data['العقود'].append(p)
        if extra is not None:
            meta=details(data,ident)
            if meta:meta.update(extra)
            else:data['تفاصيل العقود'].append({'المعرف':ident,**extra})
    elif action=='activate':
        c=lookup(data,'العقود',p['contract'])
        if lookup(data,'العملاء',c['العميل'])['الحالة']!='نشط':raise Error('العميل غير نشط.')
        if c['النوع']==MANAGED and lookup(data,'المستثمرون',c['المستثمر'])['الحالة']!='نشط':raise Error('المستثمر غير نشط.')
        if c['النوع']==EXTERNAL and lookup(data,'المكاتب',c['المكتب'])['الحالة']!='نشط':raise Error('المكتب غير نشط.')
        post(data,activation_rows(c,details(data,c['المعرف'])),p['date'],c['المعرف'],c.get('المستثمر') or '',p.get('ref'),p.get('notes'),request)
    elif action in ['receipt','settle','refund']:
        c=lookup(data,'العقود',p['contract']);a=money(p['amount']);hd=a if c['النوع']==MANAGED else 0
        kind={'receipt':'سداد','settle':'تسوية','refund':'استرداد فائض'}[action]
        post(data,[(kind,a,a,hd if action!='refund' else 0)],p['date'],c['المعرف'],c.get('المستثمر') or '',p.get('ref'),p.get('notes'),request,money(p.get('discount')))
    elif action=='cash':
        a=money(p['amount']);kind=p['kind']
        post(data,[(kind,a,a if kind=='إضافة رأس مال' else -a,-a if kind=='تحويل للمستثمر' else 0)],p['date'],investor=p.get('investor') or '',ref=p.get('ref'),notes=p.get('notes'),request=request)
    elif action in ['void','restore']:
        matches=[t for t in data['الحركات'] if t['المجموعة']==p['group']]
        if not matches:raise Error('مجموعة الحركة غير موجودة.')
        for t in matches:t['الحالة']='ملغاة' if action=='void' else 'نشطة'
    elif action=='save_followup':
        p['المعرف']=p.get('المعرف') or uid('FU');old=next((r for r in data['المتابعات'] if r['المعرف']==p['المعرف']),None)
        if old:old.update(p)
        else:data['المتابعات'].append(p)
    elif action=='amend_contract':
        cid=p['المعرف'];c=lookup(data,'العقود',cid);extra=p.pop('details',{})
        state=next(x for x in snapshot(data)['contracts'] if x['id']==cid)
        if not state['active'] or state['settled']:raise Error('تعديل الشروط يتطلب عقدًا منفذًا غير مسوى.')
        for key in ['العميل','النوع','المستثمر','المكتب']:
            if (p.get(key) or '')!=(c.get(key) or ''):raise Error('لا يمكن تغيير العميل أو مالك العقد المنفذ. استخدم إلغاء العقد عند فسخ الصفقة.')
        groups={t['المجموعة'] for t in data['الحركات'] if t.get('العقد')==cid and t['النوع']=='تنفيذ' and t['الحالة']=='نشطة'}
        oldrows=[t for t in data['الحركات'] if t['المجموعة'] in groups]
        if len(groups)!=1:raise Error('تعذر تحديد التنفيذ الأصلي.')
        first=oldrows[0]
        for t in oldrows:t['الحالة']='ملغاة'
        c.update(p);meta=details(data,cid)
        if not meta:meta={'المعرف':cid};data['تفاصيل العقود'].append(meta)
        meta.update(extra)
        group=post(data,activation_rows(c,meta),first['التاريخ'],cid,c.get('المستثمر') or '',first.get('المرجع'),'تصحيح شروط العقد والتنفيذ',request)
        for t in data['الحركات']:
            if t['المجموعة']==group:t['الوقت']=first['الوقت']
    elif action=='save_contract_info':
        c=lookup(data,'العقود',p['contract']);meta=details(data,c['المعرف'])
        if not meta:
            meta={'المعرف':c['المعرف']};data['تفاصيل العقود'].append(meta)
        meta.update({k:p[k] for k in ['نوع السلعة','وصف السلعة','قيمة سند لأمر'] if k in p})
        if 'ملاحظات' in p:c['ملاحظات']=p['ملاحظات']
    elif action=='cancel_contract':
        cid=p['contract'];lookup(data,'العقود',cid)
        if details(data,cid).get('ملغى'):raise Error('العقد ملغى بالفعل.')
        if not txt(p.get('reason')):raise Error('سبب الإلغاء مطلوب.')
        active=[t for t in data['الحركات'] if t.get('العقد')==cid and t['الحالة']=='نشطة']
        if active and not p.get('reversed'):raise Error('أكد إعادة جميع الأموال والتحويلات المرتبطة بالعقد فعليًا قبل الإلغاء.')
        for t in active:t['الحالة']='ملغاة'
        meta=details(data,cid)
        if not meta:meta={'المعرف':cid};data['تفاصيل العقود'].append(meta)
        meta.update({'ملغى':True,'سبب الإلغاء':txt(p['reason']),'تاريخ الإلغاء':today()})
    elif action=='delete_record':
        sheet=p['sheet'];ident=p['id']
        if sheet not in ['العملاء','المستثمرون','المكاتب','العقود','المتابعات']:raise Error('هذا السجل لا يقبل الحذف.')
        lookup(data,sheet,ident)
        key={'العملاء':'العميل','المستثمرون':'المستثمر','المكاتب':'المكتب','العقود':'العقد','المتابعات':None}[sheet]
        if key and any(r.get(key)==ident for other in ['العقود','الحركات','المتابعات'] if other!=sheet for r in data[other]):raise Error('السجل مرتبط بعقود أو حركات أو متابعات. استخدم الإيقاف أو الإلغاء بدل الحذف.')
        data[sheet]=[r for r in data[sheet] if r['المعرف']!=ident]
        if sheet=='العقود':data['تفاصيل العقود']=[r for r in data['تفاصيل العقود'] if r['المعرف']!=ident]
    elif action=='edit_transaction':
        ts=[t for t in data['الحركات'] if t['المجموعة']==p['group']]
        if len(ts)!=1 or ts[0]['النوع'] not in ['سداد','تسوية','استرداد فائض','إضافة رأس مال','سحب رأس مال','مصروف مكتب','تحويل للمستثمر'] or ts[0]['الحالة']!='نشطة':raise Error('مجموعة التنفيذ تُصحح بالإلغاء وإعادة التنفيذ.')
        t=ts[0];amount=money(p['amount']);t['الحالة']='ملغاة'
        bd=amount if money(t['أثر البنك'])>0 else -amount
        hd=(1 if money(t['أثر أموال المستثمر'])>0 else -1)*amount if money(t['أثر أموال المستثمر']) else 0
        post(data,[(t['النوع'],amount,bd,hd)],p['date'],t.get('العقد') or '',t.get('المستثمر') or '',p.get('ref'),str(p.get('notes') or '')+' | تصحيح '+t['المعرف'],request,money(p.get('discount',t.get('خصم التسوية'))))
    elif action=='settings':
        if not 1<=int(p['default_months'])<=120 or not 0<=float(p['default_profit_rate'])<=10:raise Error('الإعدادات خارج المدى.')
        for r in data['النظام']:
            if r['المفتاح'] in p:r['القيمة']=p[r['المفتاح']]
    else:raise Error('عملية غير معروفة.')
    for r in data['النظام']:
        if r['المفتاح']=='version':r['القيمة']=VERSION
    snapshot(data)
    return data

def read_excel(raw):
    try:wb=load_workbook(BytesIO(raw),data_only=False)
    except Exception:raise Error('الملف ليس ملف Excel صالحًا.')
    data={}
    for name,headers in HEADERS.items():
        if name=='تفاصيل العقود' and name not in wb.sheetnames:
            data[name]=[];continue
        if name not in wb.sheetnames:raise Error('هذا ليس ملف النظام الجديد V2. استخدم Installment_System_V2.xlsx.')
        ws=wb[name];actual=[x.value for x in ws[1]]
        if actual[:len(headers)]!=headers:raise Error('أعمدة ورقة '+name+' غير متوافقة مع V2.')
        if name=='الأقساط':
            data[name]=[]
            continue
        count=INPUT_COLUMNS.get(name,len(headers));data[name]=[]
        for row in ws.iter_rows(min_row=2):
            if row[0].value in [None,'']:continue
            record={headers[i]:row[i].value for i in range(count)}
            if any(row[i].data_type=='f' for i in range(count)):raise Error('حقول البيانات الأصلية لا تقبل معادلات. استخدم التطبيق لتعديلها.')
            data[name].append(record)
    if setting(data,'format')!=FORMAT:raise Error('بنية الملف ليست V2.')
    data.setdefault('تفاصيل العقود',[])
    snapshot(data);return data

def put(ws,row,col,value):
    cell=ws.cell(row,col);cell.value=value
    if isinstance(value,str):cell.data_type='s'

def write_excel(data,template=None):
    """Application serializer. Rebuild calculated sheets from the raw records."""
    snap=snapshot(data)
    wb=load_workbook(BytesIO(template)) if template else load_workbook(ROOT/'Installment_System_V2.xlsx')
    for name,headers in HEADERS.items():
        ws=wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        for j,key in enumerate(headers,1):put(ws,1,j,key)
        for key in list(ws.tables):del ws.tables[key]
        ws.delete_rows(2,ws.max_row)
        for i,record in enumerate(data[name],2):
            for j,key in enumerate(headers,1):
                if key in record:put(ws,i,j,record[key])
    schedules=[]
    for z in snap['parts']:
        schedules.append([z['id'],z['contract'],z['sequence'],z['due'],sar(z['amount']),sar(z['paid']),sar(z['waived']),sar(z['remaining']),z['status'],z['late_days'],z['paid_date']])
    ws=wb['الأقساط'];ws.delete_rows(2,ws.max_row)
    for i,record in enumerate(schedules,2):
        for j,v in enumerate(record,1):put(ws,i,j,v)
        ws.cell(i,8,0 if snap['parts'][i-2]['status']=='ملغى' else f'=MAX(0,E{i}-F{i}-G{i})')
    active={s['id']:s for s in snap['contracts']};sn=max(2,len(schedules)+1)
    ws=wb['العقود']
    for i,c in enumerate(data['العقود'],2):
        s=active[c['المعرف']];p=s['p'];managed=c['النوع']==MANAGED
        ws.cell(i,18,f'=J{i}-I{i}')
        ws.cell(i,19,f'=I{i}+N{i}' if managed else f'=I{i}+M{i}+P{i}')
        ws.cell(i,20,f'=K{i}-S{i}')
        ws.cell(i,21,f'=N{i}+O{i}-P{i}' if managed else f'=T{i}+O{i}')
        ws.cell(i,22,f'=IF(S{i}=0,0,T{i}/S{i})')
        ws.cell(i,23,f'=SUMIFS(\'الأقساط\'!$F$2:$F${sn},\'الأقساط\'!$B$2:$B${sn},A{i})+{sar(s["down"]) if s["active"] else 0}')
        ws.cell(i,24,f'=SUMIFS(\'الأقساط\'!$G$2:$G${sn},\'الأقساط\'!$B$2:$B${sn},A{i})')
        ws.cell(i,25,0 if s['cancelled'] else f'=K{i}-W{i}-X{i}')
        put(ws,i,26,s['status']);put(ws,i,27,sar(s['surplus_due']));put(ws,i,28,s['activated'])
        ws.cell(i,22).number_format='0.0%'
    bytx={t['المعرف']:t for t in snap['ledger']}
    ws=wb['الحركات']
    for i,tr in enumerate(data['الحركات'],2):
        t=bytx.get(tr['المعرف'])
        if t:
            for j,k in [(16,'bank'),(17,'held'),(18,'free')]:put(ws,i,j,sar(t[k]))
    dash=wb['الملخص'];dash.delete_rows(1,dash.max_row)
    dash.append(['إدارة التقسيط '+VERSION,'القيمة بالريال / العدد']);dash.append(['رصيد البنك',sar(snap['bank'])]);dash.append(['أموال المستثمرين المحتجزة',sar(snap['held'])]);dash.append(['السيولة الحرة للمكتب','=B2-B3'])
    funded=[s for s in snap['contracts'] if s['active']]
    for row in [('متبقي عقودك',sar(sum(s['remaining'] for s in funded if s['c']['النوع']!=MANAGED))),('متبقي عقود المستثمرين',sar(sum(s['remaining'] for s in funded if s['c']['النوع']==MANAGED))),('الربح المتوقع للمكتب',sar(sum(s['p']['office_profit'] for s in funded))),('تحصيل الأقساط',sar(sum(s['paid'] for s in funded))),('الخصومات',sar(sum(s['waived'] for s in funded))),('فائض مستحق لدى المكاتب',sar(sum(s['surplus_due'] for s in funded))),('عدد العملاء',len(data['العملاء'])),('العقود المنفذة',len(funded)),('آخر حفظ',datetime.now(TZ).isoformat(timespec='seconds'))]:dash.append(row)
    for ws in wb:
        ws.sheet_view.rightToLeft=True;ws.sheet_view.showGridLines=False;ws.freeze_panes='C2' if ws.title!='الملخص' else None
        ws.row_dimensions[1].height=38
        for cell in ws[1]:cell.font=Font(name='Arial',size=11,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1B3744');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height=27
            for cell in row:
                cell.font=Font(name='Arial',size=11,color='156D5A' if cell.data_type=='f' else '18323D')
                cell.alignment=Alignment(horizontal='right',vertical='center')
                if isinstance(cell.value,(int,float)) and cell.number_format=='General':cell.number_format='#,##0.00;(#,##0.00);"—"'
                if isinstance(cell.value,(date,datetime)):cell.number_format='yyyy-mm-dd'
        for j in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(j)].width=25 if ws.title!='الملخص' else (45 if j==1 else 35)
        if ws.title in HEADERS and ws.max_row>1:
            table=Table(displayName='Data'+str(wb.sheetnames.index(ws.title)),ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}')
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(table)
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    b=BytesIO();wb.save(b);return b.getvalue()

def digest(raw):return hashlib.sha256(raw).hexdigest()
def atomic(path,raw):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    fd,tmp=tempfile.mkstemp(prefix='.save-',suffix='.xlsx',dir=path.parent)
    try:
        with os.fdopen(fd,'wb') as f:f.write(raw);f.flush();os.fsync(f.fileno())
        os.replace(tmp,path)
    finally:
        if os.path.exists(tmp):os.unlink(tmp)

class LocalStore:
    def __init__(self,path):self.path=Path(path)
    def read(self):
        with FileLock(str(self.path)+'.lock'):
            raw=self.path.read_bytes();read_excel(raw);return raw,digest(raw)
    def commit(self,tag,action,payload,request):
        with FileLock(str(self.path)+'.lock'):
            raw=self.path.read_bytes()
            if digest(raw)!=tag:raise Conflict('تغير المصدر. اضغط تحديث ثم راجع العملية.')
            data=perform(read_excel(raw),action,payload,request);new=write_excel(data,raw)
            back=self.path.parent/'backups'/(datetime.now(TZ).strftime('%Y%m%d_%H%M%S_%f')+'.xlsx');atomic(back,raw)
            atomic(self.path,new);return new,digest(new)

class CloudStore:
    def __init__(self,token,folder,name):
        self.session=requests.Session();self.session.headers['Authorization']='Bearer '+token
        self.graph='https://graph.microsoft.com/v1.0';self.folder=folder.strip('/');self.name=name
        if name!='Installment_System_V2.xlsx':raise Error('اسم ملف النسخة الجديدة يجب أن يكون Installment_System_V2.xlsx.')
        self.path=quote(self.folder+'/'+name,safe='/');self.item=''
    def check(self,r):
        if r.status_code in [409,412]:raise Conflict('تغير المصدر أثناء الحفظ. حدّث وتحقق من سجل الحركات قبل الإعادة.')
        if r.status_code==404:raise Error('ارفع Installment_System_V2.xlsx داخل مجلد '+self.folder+' في OneDrive أولًا.')
        if r.status_code in [401,403]:raise Error('انتهت جلسة Microsoft أو لا تملك صلاحية الملف. أعد تسجيل الدخول.')
        if r.status_code==423:raise Error('الملف مقفل. أغلقه في Excel وحاول مجددًا.')
        if not 200<=r.status_code<300:raise Error('تعذر الاتصال بـOneDrive. HTTP '+str(r.status_code))
    def meta(self):
        r=self.session.get(self.graph+'/me/drive/root:/'+self.path,params={'$select':'id,eTag'},timeout=35);self.check(r)
        m=r.json();self.item=m['id']
        if not m.get('eTag'):raise Conflict('تعذر تحديد نسخة الملف.')
        return m
    def read(self):
        a=self.meta();r=self.session.get(self.graph+'/me/drive/items/'+quote(self.item,safe='')+'/content',timeout=90);self.check(r)
        b=self.meta()
        if a['eTag']!=b['eTag']:raise Conflict('تغير الملف أثناء القراءة. حدّث مرة أخرى.')
        read_excel(r.content);return r.content,b['eTag']
    def commit(self,tag,action,payload,request):
        raw,current=self.read();data=read_excel(raw)
        if any(t['رقم الطلب']==request for t in data['الحركات']):return raw,current
        if current!=tag:raise Conflict('تغير Excel بعد فتح الشاشة. اضغط تحديث وراجع البيانات ثم احفظ.')
        new=write_excel(perform(data,action,payload,request),raw)
        # Keep recoverable bytes before the conditional remote write.
        backup=Path(tempfile.gettempdir())/'installments_v2_recovery'/digest(self.item.encode())/(datetime.now(TZ).strftime('%Y%m%d_%H%M%S_%f')+'.xlsx');atomic(backup,raw)
        r=self.session.post(self.graph+'/me/drive/items/'+quote(self.item,safe='')+'/createUploadSession',headers={'If-Match':current},json={'item':{'@microsoft.graph.conflictBehavior':'replace'}},timeout=45);self.check(r)
        url=r.json()['uploadUrl']
        if self.meta()['eTag']!=current:raise Conflict('تغير الملف أثناء تجهيز الحفظ. حدّث المصدر.')
        try:r=requests.put(url,data=new,headers={'Content-Length':str(len(new)),'Content-Range':f'bytes 0-{len(new)-1}/{len(new)}'},timeout=120)
        except requests.RequestException:raise Conflict('لم يصل تأكيد الحفظ. حدّث المصدر وتحقق من رقم الطلب قبل إعادة العملية.')
        self.check(r)
        if r.status_code not in [200,201] or not r.json().get('eTag'):raise Conflict('لم يصل تأكيد نهائي. حدّث المصدر وتحقق من رقم الطلب.')
        return new,r.json()['eTag']

def config(section,key,default=''):
    try:return st.secrets[section].get(key,default)
    except Exception:return default
EN_TEXT = {
  "إدارة التقسيط": "Installment Management",
  "إدارة التقسيط | إبراهيم": "Installment Management | Ibrahim",
  "إبراهيم · الإصدار ": "Ibrahim · Version ",
  "الملخص": "Overview",
  "العملاء": "Customers",
  "المستثمرون": "Investors",
  "المكاتب": "External Offices",
  "العقود والتحصيل": "Contracts & Collections",
  "الخزينة والحركات": "Treasury & Transactions",
  "متابعة العملاء": "Customer Follow-ups",
  "التقارير": "Reports",
  "الإعدادات والنسخ": "Settings & Backups",
  "الصفحات": "Pages",
  "التنبيهات والمتابعة": "Alerts & Follow-ups",
  "عقود متأخرة": "Overdue contracts",
  "متابعات مستحقة": "Follow-ups due",
  "إجمالي المتأخر": "Total overdue",
  "فتح عقد من التنبيهات": "Contract from alerts",
  "فتح العقد المحدد": "Open selected contract",
  "المستحق والمتأخر ومتابعات اليوم": "Due installments and follow-ups",
  "الملخص المالي": "Financial Summary",
  "السيولة الحرة": "Available office cash",
  "رصيد البنك": "Bank balance",
  "أمانات المستثمرين": "Investor funds held",
  "ربح المكتب المتوقع": "Expected office profit",
  "متبقي عقودك": "Your receivables",
  "متبقي عقود المستثمرين": "Investor receivables",
  "العقود والإجراءات السريعة": "Contracts & Quick Actions",
  "إضافة عميل": "Add customer",
  "عقد جديد": "New contract",
  "تسجيل سداد": "Record payment",
  "الإجراء": "Action",
  "عرض / تعديل عقد": "View / edit contract",
  "بحث بالعميل أو رقم العقد أو السلعة": "Search customer, contract ID or item",
  "حالة العقد": "Contract status",
  "تصفية بالمستثمر": "Filter by investor",
  "اختر العقد لعرضه أو تعديله": "Select a contract to view or edit",
  "الكل": "All",
  "مسودة": "Draft",
  "منتظم": "Current",
  "متأخر": "Overdue",
  "مكتمل": "Completed",
  "مسدد مبكرًا": "Settled early",
  "ملغى": "Cancelled",
  "مسدد": "Paid",
  "مستحق اليوم": "Due today",
  "قادم": "Upcoming",
  "سداد جزئي": "Partly paid",
  "لم ينفذ": "Not activated",
  "نشط": "Active",
  "موقوف": "Suspended",
  "مؤرشف": "Archived",
  "نشطة": "Active",
  "ملغاة": "Voided",
  "مفتوحة": "Open",
  "مكتملة": "Completed",
  "غير مصنف": "Unrated",
  "غير مسجل": "Not recorded",
  "بدون": "None",
  "إضافة أو تعديل": "Add or edit",
  "إضافة سجل جديد": "Add new record",
  "الاسم": "Full name",
  "الجوال": "Mobile number",
  "الهوية": "ID number",
  "الحالة": "Status",
  "المستوى": "Rating",
  "الحد الائتماني (صفر = غير محدد)": "Credit limit (0 = unlimited)",
  "العنوان": "Address",
  "جهة العمل": "Employer",
  "الكفيل": "Guarantor",
  "الدخل الشهري": "Monthly income",
  "جوال الكفيل": "Guarantor mobile",
  "ملاحظات": "Notes",
  "حفظ في Excel": "Save to Excel",
  "حذف السجل نهائيًا": "Delete record permanently",
  "أؤكد حذف هذا السجل": "I confirm deletion",
  "حذف السجل": "Delete record",
  "إجمالي رأس المال الممول": "Total purchase funding",
  "رأس المال الممول": "Purchase funding",
  "تكلفة السلع المشتراة": "Cost of purchased goods",
  "أمانات نقدية متاحة": "Investor cash held",
  "أقساط متبقية للمستثمر": "Remaining investor receivables",
  "أموال محتجزة": "Funds held",
  "متبقي العقود": "Outstanding receivables",
  "نوع الصفقة": "Transaction model",
  "تمويل ذاتي": "Self-funded",
  "استثمار عبر مكتب": "Investing through an office",
  "إدارة تمويل مستثمر": "Managing an investor contract",
  "العميل": "Customer",
  "المستثمر": "Investor",
  "المكتب المنفذ": "Executing office",
  "تكلفة الشراء الفعلية": "Actual purchase cost",
  "سعر الكاش / أصل التمويل للعميل": "Reference cash price / customer principal",
  "سعر البيع الإجمالي شامل الربح والدفعة الأولى": "Total sale price including profit and down payment",
  "التمويل المحول لشراء السلعة (دون الرسوم)": "Purchase funding transferred (excluding fees)",
  "رسوم مدفوعة للمكتب المنفذ": "Fees paid to executing office",
  "رسومك الثابتة على المستثمر": "Your fixed fee charged to investor",
  "رسوم إضافية يدفعها العميل خارج الأقساط": "Additional customer fees outside sale price",
  "مصاريف أخرى تتحملها أنت": "Other expenses borne by your office",
  "عدد الأقساط الشهرية": "Number of monthly installments",
  "تاريخ العقد": "Contract date",
  "أول استحقاق": "First installment due",
  "نوع السلعة": "Item type",
  "وصف السلعة": "Item description",
  "هل توجد دفعة أولى؟": "Is there a down payment?",
  "قيمة الدفعة الأولى": "Down payment amount",
  "قيمة سند لأمر": "Promissory note amount",
  "ملاحظات العقد": "Contract notes",
  "فرق الشراء": "Purchase price difference",
  "ربح المستثمر المتوقع": "Expected investor profit",
  "القسط التقريبي": "Approximate installment",
  "حفظ مسودة العقد": "Save draft contract",
  "حفظ تصحيح شروط العقد": "Save corrected contract terms",
  "راجعت شروط العقد الجديدة وهي تصحيح للصفقة والتحويلات الفعلية": "I reviewed the corrected terms and confirmed the actual transaction amounts",
  "اختر العقد": "Select contract",
  "إجمالي التقسيط": "Total sale price",
  "المحصل": "Collected",
  "المتبقي": "Outstanding",
  "قيود التنفيذ التي ستسجل:": "Entries to be posted on activation:",
  "تاريخ التنفيذ": "Activation date",
  "مرجع التحويل / العملية": "Transfer / transaction reference",
  "ملاحظات التنفيذ": "Activation notes",
  "تمت جميع التحويلات والشراء وتحصيل الرسوم الموضحة أعلاه فعليًا": "All listed transfers, purchases and fees have actually been completed",
  "تنفيذ الصفقة": "Activate contract",
  "تعديل المسودة": "Edit draft",
  "العملية": "Operation",
  "سداد قسط أو مبلغ جزئي": "Installment or partial payment",
  "تسوية مبكرة": "Early settlement",
  "المبلغ المستلم فعليًا": "Amount actually received",
  "خصم التسوية المتفق عليه (للتسوية المبكرة فقط)": "Agreed discount (early settlement only)",
  "تاريخ السداد": "Payment date",
  "مرجع السداد": "Payment reference",
  "المحصل / ملاحظات": "Collector / notes",
  "تسجيل السداد": "Save payment",
  "المبلغ المسترد": "Refund received",
  "تاريخ استرداد الفائض": "Surplus refund date",
  "مرجع الاسترداد": "Refund reference",
  "تسجيل استرداد الفائض": "Record surplus refund",
  "تعديل الأسعار والدفعة الأولى والجدولة": "Correct prices, down payment and schedule",
  "تعديل بيانات السلعة والسند والملاحظات": "Edit item details, note amount and notes",
  "حفظ تعديلات العقد": "Save contract changes",
  "إلغاء العقد": "Cancel contract",
  "سبب إلغاء العقد": "Reason for cancellation",
  "أؤكد إعادة السلعة وتسوية وإرجاع جميع مبالغ هذا العقد فعليًا، وأريد إلغاء قيوده": "I confirm the goods and all related funds have actually been returned or settled, and request reversal of the entries",
  "تأكيد إلغاء العقد": "Confirm cancellation",
  "البنك": "Bank",
  "المتاح للمكتب": "Available to office",
  "نوع الحركة": "Transaction type",
  "إضافة رأس مال": "Capital contribution",
  "سحب رأس مال": "Capital withdrawal",
  "مصروف مكتب": "Office expense",
  "تحويل للمستثمر": "Transfer to investor",
  "المبلغ": "Amount",
  "تاريخ الحركة": "Transaction date",
  "المرجع": "Reference",
  "الوصف": "Description",
  "حفظ الحركة": "Save transaction",
  "سجل الحركات": "Transaction history",
  "تعديل حركة مالية": "Correct financial transaction",
  "اختر حركة للتعديل": "Select transaction to correct",
  "المبلغ المصحح": "Corrected amount",
  "تاريخ الحركة المصحح": "Corrected transaction date",
  "المرجع المصحح": "Corrected reference",
  "وصف التصحيح": "Correction notes",
  "خصم التسوية المصحح": "Corrected settlement discount",
  "حفظ تصحيح الحركة": "Save correction",
  "تصحيح حركة بإلغاء المجموعة أو استعادتها": "Void or restore a posting group",
  "المجموعة": "Posting group",
  "راجعت المجموعة كاملة وأريد تغيير حالتها": "I reviewed the whole posting group and want to change its status",
  "إلغاء المجموعة": "Void group",
  "استعادة المجموعة": "Restore group",
  "إضافة أو تعديل متابعة": "Add or edit follow-up",
  "العقد (اختياري)": "Contract (optional)",
  "تاريخ التواصل": "Contact date",
  "وسيلة التواصل": "Contact method",
  "اتصال": "Phone call",
  "واتساب": "WhatsApp",
  "زيارة": "Visit",
  "رسالة": "Message",
  "أخرى": "Other",
  "نتيجة التواصل": "Contact outcome",
  "يوجد وعد سداد": "Payment promise received",
  "موعد الوعد": "Promised payment date",
  "المبلغ الموعود": "Promised amount",
  "موعد المتابعة القادمة": "Next follow-up date",
  "الملاحظات": "Notes",
  "حفظ المتابعة": "Save follow-up",
  "من تاريخ": "From date",
  "إلى تاريخ": "To date",
  "تحصيل أقساطك خلال الفترة": "Your collections during period",
  "تحصيل للمستثمرين": "Collections for investors",
  "رسوم مكتسبة": "Fees earned",
  "صافي أثر البنك": "Net change in bank balance",
  "قائمة التعثر الحالية": "Current overdue contracts",
  "حقوق المستثمرين الحالية": "Current investor balances",
  "أمانات نقدية": "Cash held",
  "أقساط غير محصلة": "Uncollected installments",
  "تنزيل نسخة Excel الحالية": "Download current Excel backup",
  "عدد الأقساط الافتراضي": "Default monthly installments",
  "الربح الإجمالي المقترح %": "Suggested total markup %",
  "حفظ الإعدادات": "Save settings",
  "الحفظ المحلي في Excel": "Local Excel storage",
  "متصل بملف OneDrive V2": "Connected to OneDrive V2 workbook",
  "تسجيل الخروج": "Sign out",
  "↻ تحديث من Excel": "↻ Refresh from Excel",
  "آخر حفظ بهذه الجلسة: ": "Last saved in this session: ",
  "لم يُحفظ بعد": "Not saved yet",
  "تم الحفظ في Excel بنجاح.": "Saved to Excel successfully.",
  "لا توجد سجلات حتى الآن.": "No records yet.",
  "البيان": "Entry",
  "أثر البنك": "Bank impact",
  "أثر أمانات المستثمر": "Investor funds impact",
  "المعرف": "ID",
  "العقد": "Contract",
  "النوع": "Type",
  "المكتب": "Office",
  "تكلفة الشراء": "Purchase cost",
  "سعر الكاش": "Cash price",
  "رسوم مكتب مدفوعة": "Office fees paid",
  "رسوم مكتب ثابتة": "Fixed office fee",
  "رسوم العميل": "Customer fees",
  "مصاريف أخرى": "Other expenses",
  "تكلفة المستثمر": "Investor cost",
  "عائد المستثمر": "Investor return",
  "الخصم": "Discount",
  "فائض لدى المكتب": "Surplus held by office",
  "الاستحقاق القادم": "Next due date",
  "التسلسل": "Sequence",
  "الاستحقاق": "Due date",
  "القسط": "Installment",
  "أيام التأخير": "Days late",
  "التمويل المحول": "Funding transferred",
  "التاريخ": "Date",
  "الوقت": "Time",
  "أثر أموال المستثمر": "Investor funds impact",
  "خصم التسوية": "Settlement discount",
  "رقم الطلب": "Request ID",
  "الرصيد البنكي": "Bank balance",
  "أموال المستثمرين": "Investor funds",
  "المتابعة": "Follow-up",
  "وعد السداد": "Payment promise",
  "النتيجة": "Outcome",
  "المتابعة القادمة": "Next follow-up",
  "سداد": "Payment",
  "تسوية": "Settlement",
  "تنفيذ": "Activation",
  "دفعة أولى": "Down payment",
  "شراء للمستثمر": "Investor purchase",
  "رسوم ثابتة": "Fixed fee",
  "رسوم عميل": "Customer fee",
  "مصروف صفقة": "Transaction expense",
  "استرداد فائض": "Surplus refund",
  "عقد ملغى: ": "Cancelled contract: ",
  "السلعة: ": "Item: ",
  " · سعر الشراء: ": " · Purchase cost: ",
  " · سعر البيع الإجمالي: ": " · Total sale price: ",
  " · الدفعة الأولى: ": " · Down payment: ",
  " · سند لأمر: ": " · Promissory note: ",
  "القيمة المقترحة وفق الإعدادات: ": "Suggested value from settings: ",
  " ريال. عدّل إجمالي التقسيط حسب الاتفاق.": " SAR. Adjust the total sale price to match the agreement.",
  "فائض لدى المكتب يجب استرداده وتسجيله: ": "Surplus held by the office, to recover and record: ",
  " ريال. الربح المتوقع يفترض استرداده.": " SAR. Expected profit assumes this surplus is recovered.",
  "المطلوب استلامه من المستثمر عند التنفيذ: ": "Amount due from the investor on activation: ",
  " ريال. أمانته المتبقية بعد الشراء: ": " SAR. Investor cash remaining after purchase: ",
  "المتاح تحويله للمستثمر: ": "Available for transfer to investor: ",
  "فائض مستحق لك لدى المكتب: ": "Surplus due to you from the office: ",
  "عميل المتابعة: ": "Follow-up customer: ",
  "تعذر فتح النظام: ": "Unable to open the application: ",
  "رقم الطلب: ": "Request ID: ",
  ". عند انقطاع الاتصال، حدّث المصدر قبل إعادة الحفظ.": ". If disconnected, refresh the source before retrying.",
  "رقم الهوية يجب أن يكون 10 أرقام بالضبط.": "ID number must contain exactly 10 digits.",
  "رقم الجوال يجب أن يكون 10 أرقام ويبدأ بـ05، مثال 0512345678.": "Mobile must contain 10 digits and start with 05, e.g. 0512345678.",
  "جوال الكفيل يجب أن يكون 10 أرقام ويبدأ بـ05.": "Guarantor mobile must contain 10 digits and start with 05.",
  "أدخل نوع السلعة.": "Enter the item type.",
  "نوع السلعة مطلوب.": "Item type is required.",
  "أدخل دفعة أولى أكبر من صفر أو ألغِ اختيار وجودها.": "Enter a positive down payment or uncheck the down payment option.",
  "أكد إتمام العمليات الفعلية أولًا.": "Confirm that the actual transactions are completed first.",
  "أضف واختر العميل والجهة الممولة أولًا.": "Add and select the customer and funding party first.",
  "الحذف للسجلات غير المرتبطة فقط. الحركات المالية تُصحح بالإلغاء للحفاظ على سجلها.": "Only unlinked records can be deleted. Financial transactions are voided to preserve their history.",
  "رأس المال الممول هو مجموع تمويل العقود المنفذة دون رسوم المكتب؛ رقم تراكمي قد يشمل إعادة تدوير المال، ولا يساوي الرصيد النقدي الحالي.": "Total purchase funding excludes office fees. This cumulative figure may include recycled capital and is not the current cash balance.",
  "الربح المتوقع يفترض اكتمال السداد؛ أمانات المستثمرين مستبعدة من السيولة الحرة.": "Expected profit assumes full collection. Investor funds are excluded from available office cash.",
  "الحسابات تتحدث عند تعديل الحقول. حفظ العقد ينشئ مسودة؛ التنفيذ يسجل النقد والرسوم الفعلية.": "Calculations update as fields change. Saving creates a draft; activation records actual cash and fees.",
  "كامل تحصيلات الأقساط للمستثمر. رسومك الثابتة يدفعها المستثمر إضافة إلى مبلغ التمويل عند التنفيذ. فرق الشراء يبقى أمانة للمستثمر.": "All customer collections belong to the investor. The investor pays your fixed fee in addition to purchase funding. Purchase savings remain the investor's funds.",
  "كامل مبلغ السداد يضاف إلى أموال المستثمر المحتجزة.": "The entire payment is added to the investor funds held.",
  "يحفظ التنفيذ السابق كملغى ويعيد حساب جميع الحركات. يرفض التصحيح إذا تعارض مع التحصيلات أو حقوق المستثمرين. لا يغيّر العقد الموقع خارج التطبيق.": "The previous activation is retained as void and all transactions are recalculated. Corrections conflicting with collections or investor balances are rejected. Signed documents outside the app are not changed.",
  "يمكن تعديل الأسعار والجدولة من قسم تصحيح شروط العقد أعلاه للعقود غير المسواة.": "For unsettled contracts, correct prices and the schedule in the contract terms section above.",
  "الإلغاء يحتفظ بالعقد وسجله ويُلغي حركاته محاسبيًا؛ لا ينفذ تحويلات بنكية. يرفض إذا نتج رصيد سالب أو تعارض مع حقوق المستثمرين.": "Cancellation retains the contract and its history and reverses its accounting entries. It does not make bank transfers. Negative balances or conflicts with investor funds are rejected.",
  "إلغاء التنفيذ يلغي قيوده معًا. إذا ترتب عليه سداد أو تحويل لاحق غير صالح فسيرفض النظام الإلغاء حتى تصحح الحركات التابعة.": "Voiding activation reverses its entire posting group. If subsequent payments or transfers become invalid, correct those transactions first.",
  "تحصيل أقساطك يشمل رأس المال والربح. الرسوم المكتسبة منفصلة. صافي أثر البنك يشمل التمويل ورأس المال والمصروفات والتحويلات.": "Your collections include principal and profit. Fees earned are separate. Net bank change includes funding, capital, expenses and transfers.",
  "في العقد المدار: المستثمر يملك كامل الأقساط. رسومه الثابتة تدفع إضافة إلى التمويل وتتحول لملكية المكتب عند التنفيذ. فائض التمويل بعد الشراء يبقى أمانة للمستثمر حتى تحويله.": "In managed contracts, all installments belong to the investor. The fixed fee is paid in addition to funding and becomes office income on activation. Purchase funding surplus remains the investor's funds until transferred.",
  "التصنيف A–D يدوي. فرق الشراء = سعر الكاش − التكلفة الفعلية. الربح المتوقع يحسب مرة واحدة من إجمالي التقسيط ناقص التكلفة والرسوم والمصاريف. الربح لا يضاف إلى سعر التقسيط مرة أخرى.": "Ratings A–D are manual. Purchase difference = cash price minus actual cost. Expected profit is calculated once from the sale price less costs, fees and expenses. Do not add profit to the sale price again.",
  "تعديل Excel خارج التطبيق متاح في حقول البيانات الأصلية وفق العناوين. الأقساط والملخص والحقول المحسوبة يعاد توليدها عند الحفظ. لتغيير الشروط المالية لعقد منفذ، صحح الحركات التابعة ثم ألغِ مجموعة التنفيذ أولًا.": "Only original data fields may be edited directly in Excel. Installments, summary and calculated fields are rebuilt on save. Financial corrections must remain consistent with dependent transactions.",
  "التطبيق يبدأ برصيد صفر. أضف رأس المال الحقيقي من صفحة الخزينة. ملف V1 مستقل ولا يُستبدل بهذا الملف.": "The application starts at zero. Add actual capital in Treasury. The V1 workbook is separate.",
  "إلزامي: 10 أرقام تبدأ بـ05، مثال 0512345678.": "Required: 10 digits starting with 05, e.g. 0512345678.",
  "10 أرقام بالضبط؛ إلزامي للعميل. تُقبل الأرقام العربية والإنجليزية.": "Exactly 10 digits; required for customers. Arabic and Western digits are accepted.",
  "اختياري؛ عند إدخاله يجب أن يكون 10 أرقام تبدأ بـ05.": "Optional; if entered, use 10 digits starting with 05.",
  "الاسم الكامل كما يظهر في المستندات.": "Full name as shown on the documents.",
  "10 أرقام بالضبط دون مسافات؛ إلزامي للعميل. للمستثمر اختياري حتى تتوفر هويته.": "Exactly 10 digits without spaces. Required for customers; optional for investors.",
  "المبلغ المدفوع فعليًا للمورّد بالريال؛ أساس حساب تكلفة الصفقة.": "Actual amount paid to the supplier in SAR; used to calculate transaction cost.",
  "سعر البيع النقدي المرجعي؛ قد يختلف عن تكلفة شرائك الفعلية.": "Reference cash sale price; may differ from your actual purchase cost.",
  "سعر البيع الإجمالي شامل الدفعة الأولى والربح، دون رسوم العميل الإضافية. لا تضف الربح إليه مرة أخرى.": "Total sale price includes the down payment and profit, but excludes extra customer fees. Do not add profit again.",
  "المبلغ المحول للشراء فقط؛ فائضه عن التكلفة يبقى حقًا لصاحب التمويل.": "Purchase funding only. Any surplus over cost belongs to the funding owner.",
  "رسوم تدفعها أنت للمكتب الخارجي؛ تزيد تكلفتك وتقلل ربحك.": "Fees you pay to the external office; these increase cost and reduce profit.",
  "رسم مستلم من المستثمر إضافة إلى تمويل الشراء؛ دخل للمكتب وحده.": "A fee received from the investor in addition to purchase funding; office income only.",
  "رسم مستقل للمكتب يُحصّل عند التنفيذ ولا يدخل ضمن سعر البيع أو سند لأمر تلقائيًا.": "Separate office fee collected on activation; not automatically included in the sale price or promissory note.",
  "تكاليف الصفقة بالريال يتحملها المكتب، بما في ذلك العقود المدارة.": "Transaction expenses in SAR borne by your office, including managed contracts.",
  "عدد الدفعات الشهرية بعد الدفعة الأولى؛ من 1 إلى 120.": "Number of monthly installments after the down payment; 1 to 120.",
  "تاريخ إبرام العقد؛ لا يتجاوز اليوم.": "Contract signing date; cannot be in the future.",
  "موعد أول قسط شهري، لا يسبق تاريخ العقد.": "First monthly installment date; cannot precede the contract date.",
  "قيمة السند الفعلي بالريال؛ للتوثيق ولا تُضاف إلى المديونية. صفر يعني لم تُسجل.": "Actual promissory note amount in SAR, for records only. Not added to debt. Zero means not recorded.",
  "المبلغ الذي دفعه العميل فعلًا؛ لا يتجاوز المتبقي، ويوزع على أقدم الأقساط.": "Actual customer payment, no more than outstanding; allocated to the oldest installments first.",
  "في التسوية: المبلغ المستلم + الخصم يجب أن يساويا كامل المتبقي.": "For settlement: amount received plus discount must equal the full outstanding balance.",
  "الحد الأقصى للتعرض حسب أصل التمويل. صفر يلغي تطبيق الحد.": "Maximum exposure based on financed principal. Zero disables this limit.",
  "تصنيف داخلي يدوي من A إلى D؛ لا يمثل فحصًا ائتمانيًا خارجيًا.": "Manual internal rating A to D; not an external credit check.",
  "نشط يسمح بعقود جديدة؛ موقوف يمنعها؛ مؤرشف للسجلات بلا عقود قائمة.": "Active permits new contracts; suspended prevents them; archive records without active contracts.",
  "رقم العملية البنكية أو الإيصال؛ يجب ألا يتكرر مع عملية مالية أخرى.": "Bank transaction or receipt reference; must not duplicate another financial entry.",
  "رقم العملية أو مرجع مستند الشراء لتمييزها ومنع التكرار.": "Transaction or purchase document reference to identify the entry and prevent duplicates.",
  "تاريخ التزام العميل بالسداد؛ يظهر التنبيه إذا تجاوزه دون إغلاق المتابعة.": "Promised payment date; an alert appears if overdue and the follow-up remains open.",
  "الموعد الذي يجب التواصل فيه مع العميل مرة أخرى.": "Date to contact the customer again.",
  "يظهر في تنبيهات الملخص عند حلول الموعد إذا بقيت المتابعة مفتوحة.": "Appears in Overview alerts when due if the follow-up remains open.",
  "المبلغ الفعلي بالريال؛ يجب أن يكون أكبر من صفر لحفظ حركة مالية.": "Actual amount in SAR; must be greater than zero to post a financial transaction.",
  "سبب الفسخ أو تصحيح القيد. يحتفظ النظام به مع العقد الملغى.": "Reason for cancellation or correction; retained with the cancelled contract.",
  "اختر التاريخ الصحيح للعملية أو المتابعة.": "Choose the correct transaction or follow-up date.",
  "اكتب تفاصيل تساعدك على مراجعة السجل لاحقًا؛ لا تدخل معلومات سرية غير مطلوبة.": "Add useful information for reviewing this record later; avoid unnecessary confidential information.",
  "أدخل القيمة الفعلية بالريال السعودي دون فواصل؛ تُحسب المبالغ بدقة الهللة.": "Enter the actual amount in SAR without separators. Amounts are calculated to the nearest halala.",
  "مثال: هاتف، أجهزة منزلية، سيارة. مطلوب للعقود الجديدة.": "Example: phone, home appliances, car. Required for new contracts.",
  "الموديل أو الرقم التسلسلي أو وصف يميز السلعة.": "Model, serial number or a description identifying the item.",
  "جزء من سعر البيع الإجمالي يدفعه العميل عند التنفيذ، وليس رسمًا إضافيًا.": "Part of the total sale price paid by the customer on activation, not an extra fee.",
  "تُحصّل عند تنفيذ العقد وتُطرح من المبلغ المجدول على الأقساط.": "Received on activation and deducted from the amount scheduled in installments.",
  "القيمة المثبتة في السند الفعلي؛ لا تُضاف إلى سعر البيع ولا إلى أرباح العقد. صفر يعني غير مسجل.": "Amount shown on the actual promissory note; not added to the sale price or profit. Zero means not recorded.",
  "حدد ": "Select ",
  " لهذا السجل؛ تحقق من الاختيار قبل الحفظ.": " for this record; check your selection before saving."
}

EN_TEXT.update({
  "العقد ملغى.": "The contract is cancelled.",
  "العقد ملغى بالفعل.": "The contract is already cancelled.",
  "سبب الإلغاء مطلوب.": "A reason for cancellation is required.",
  "العقد ملغى؛ لا يمكن تنفيذه.": "A cancelled contract cannot be activated.",
  "العقد الملغى لا يقبل حركات مالية نشطة.": "A cancelled contract cannot have active financial transactions.",
  "أدخل مبلغًا رقميًا صالحًا.": "Enter a valid numeric amount.",
  "التاريخ غير صالح.": "The date is invalid.",
  "النص غير صالح أو طويل جدًا.": "The text is invalid or too long.",
  "السجل المرتبط غير موجود: ": "Linked record not found: ",
  "معرف مفقود أو مكرر في ": "Missing or duplicate ID in ",
  "بيانات غير صالحة في ": "Invalid data in ",
  "الهوية مكررة في ": "Duplicate ID number in ",
  "تصنيف أو بيانات مالية للعميل غير صالحة.": "The customer rating or financial details are invalid.",
  "اختر نوع الصفقة.": "Select a transaction model.",
  "تحقق من تكلفة الشراء وسعر الكاش وإجمالي التقسيط.": "Check the purchase cost, cash price and total sale price.",
  "الرسوم والمصاريف لا تقبل السالب.": "Fees and expenses cannot be negative.",
  "التمويل الذاتي يساوي تكلفة الشراء الفعلية.": "Self-funding must equal the actual purchase cost.",
  "المبلغ المحول أقل من تكلفة الشراء.": "The transferred funding is less than the purchase cost.",
  "عقد المستثمر يستخدم رسوم المكتب الثابتة.": "Managed investor contracts use the fixed office fee.",
  "المستثمر والرسوم الثابتة يخصان العقود المدارة فقط.": "The investor and fixed fee fields apply to managed contracts only.",
  "المكتب الخارجي يخص الاستثمار عبر مكتب.": "An external office applies only to investing through an office.",
  "مدة العقد من شهر إلى 120 شهرًا.": "The contract must have 1 to 120 monthly installments.",
  "تواريخ العقد غير صالحة.": "The contract dates are invalid.",
  "تفاصيل عقد مكررة أو مرتبطة بعقد مفقود.": "Contract details are duplicated or refer to a missing contract.",
  "الدفعة الأولى من صفر إلى أقل من سعر البيع، وقيمة السند غير سالبة.": "The down payment must be zero or less than the sale price; the note amount cannot be negative.",
  "المبلغ بعد الدفعة الأولى أقل من عدد الأقساط بالهللات.": "The balance after the down payment is too small for the installment count.",
  "حالة حركة غير صالحة.": "Invalid transaction status.",
  "تاريخ الحركة غير صالح.": "Invalid transaction date.",
  "مجموعة الحركة ورقم الطلب مطلوبان.": "A posting group and request ID are required.",
  "رقم طلب مكرر.": "Duplicate request ID.",
  "مجموعة حركة غير مكتملة.": "Incomplete posting group.",
  "حركة مرتبطة بعقد مفقود.": "A transaction refers to a missing contract.",
  "الحركة تسبق العقد.": "The transaction date precedes the contract date.",
  "تنفيذ مكرر للعقد.": "The contract has already been activated.",
  "قيود التنفيذ لا تطابق شروط العقد.": "The activation entries do not match the contract terms.",
  "التنفيذ يتجاوز الحد الائتماني للعميل.": "Activation would exceed the customer's credit limit.",
  "مجموعة مالية غير معروفة.": "Unknown posting group.",
  "مبلغ الحركة لا يقبل السالب.": "A transaction amount cannot be negative.",
  "السداد يتطلب عقدًا منفذًا وغير مسوى.": "Payments require an active contract that has not been settled.",
  "مبلغ السداد أو الخصم لا يطابق المتبقي.": "The payment or discount does not match the outstanding balance.",
  "ملكية السداد غير مطابقة للعقد.": "The payment owner does not match the contract.",
  "قيد الدفعة الأولى غير مطابق.": "The down payment entry does not match the contract.",
  "استرداد الفائض يتطلب استثمارًا منفذًا عبر مكتب.": "Surplus refunds require an active investment through an office.",
  "استرداد الفائض يتجاوز المستحق أو قيده غير مطابق.": "The surplus refund exceeds the amount due or its entry is invalid.",
  "تحويل المستثمر غير صالح.": "Invalid investor transfer.",
  "قيد السيولة غير مطابق.": "The cash entry does not match its transaction type.",
  "نوع حركة غير معروف.": "Unknown transaction type.",
  "الخصم يخص التسوية فقط.": "A discount applies only to settlement.",
  "مرجع مالي مكرر.": "Duplicate financial reference.",
  "مستثمر الحركة غير موجود.": "The transaction's investor does not exist.",
  "أثر المستثمر دون معرف مستثمر.": "An investor funds impact requires an investor ID.",
  "رصيد غير كافٍ أو العملية تستخدم أموال المستثمر لصالح المكتب.": "Insufficient funds, or the transaction would use investor funds for the office.",
  "متابعة مرتبطة بعميل مختلف.": "The follow-up refers to a different customer.",
  "متابعة غير صالحة.": "Invalid follow-up.",
  "اختر تاريخًا لا يتجاوز اليوم.": "Select a date no later than today.",
  "السجل لديه عقد قائم. استخدم موقوف بدل الأرشفة.": "This record has an active contract. Suspend it instead of archiving it.",
  "شروط العقد المنفذ مقفلة. ألغِ التنفيذ والحركات التابعة أولًا.": "Executed terms are locked. Use the contract correction workflow or reverse dependent entries first.",
  "العميل غير نشط.": "The customer is not active.",
  "المستثمر غير نشط.": "The investor is not active.",
  "المكتب غير نشط.": "The office is not active.",
  "مجموعة الحركة غير موجودة.": "Posting group not found.",
  "تعديل الشروط يتطلب عقدًا منفذًا غير مسوى.": "Correcting terms requires an active, unsettled contract.",
  "لا يمكن تغيير العميل أو مالك العقد المنفذ. استخدم إلغاء العقد عند فسخ الصفقة.": "The customer or funding owner of an active contract cannot be changed. Cancel the contract if the sale was rescinded.",
  "تعذر تحديد التنفيذ الأصلي.": "The original activation could not be identified.",
  "أكد إعادة جميع الأموال والتحويلات المرتبطة بالعقد فعليًا قبل الإلغاء.": "Confirm that all related funds and transfers have actually been returned before cancellation.",
  "هذا السجل لا يقبل الحذف.": "This record cannot be deleted.",
  "السجل مرتبط بعقود أو حركات أو متابعات. استخدم الإيقاف أو الإلغاء بدل الحذف.": "This record has linked contracts, transactions or follow-ups. Suspend or cancel it instead of deleting it.",
  "مجموعة التنفيذ تُصحح بالإلغاء وإعادة التنفيذ.": "Correct an activation posting group by reversing and activating it again.",
  "الإعدادات خارج المدى.": "Settings are outside the permitted range.",
  "عملية غير معروفة.": "Unknown operation.",
  "الملف ليس ملف Excel صالحًا.": "The file is not a valid Excel workbook.",
  "هذا ليس ملف النظام الجديد V2. استخدم Installment_System_V2.xlsx.": "This is not a V2 system workbook. Use Installment_System_V2.xlsx.",
  "أعمدة ورقة ": "Worksheet columns in ",
  " غير متوافقة مع V2.": " are not compatible with V2.",
  "حقول البيانات الأصلية لا تقبل معادلات. استخدم التطبيق لتعديلها.": "Original data fields cannot contain formulas. Edit them using the app.",
  "بنية الملف ليست V2.": "The workbook structure is not V2.",
  "تغير المصدر. اضغط تحديث ثم راجع العملية.": "The source changed. Refresh and review the transaction.",
  "اسم ملف النسخة الجديدة يجب أن يكون Installment_System_V2.xlsx.": "The V2 workbook must be named Installment_System_V2.xlsx.",
  "تغير المصدر أثناء الحفظ. حدّث وتحقق من سجل الحركات قبل الإعادة.": "The source changed while saving. Refresh and check the transaction history before retrying.",
  "ارفع Installment_System_V2.xlsx داخل مجلد ": "Upload Installment_System_V2.xlsx into folder ",
  " في OneDrive أولًا.": " in OneDrive first.",
  "انتهت جلسة Microsoft أو لا تملك صلاحية الملف. أعد تسجيل الدخول.": "Your Microsoft session expired or access is denied. Sign in again.",
  "الملف مقفل. أغلقه في Excel وحاول مجددًا.": "The workbook is locked. Close it in Excel and try again.",
  "تعذر الاتصال بـOneDrive. HTTP ": "Unable to connect to OneDrive. HTTP ",
  "تعذر تحديد نسخة الملف.": "Unable to identify the workbook version.",
  "تغير الملف أثناء القراءة. حدّث مرة أخرى.": "The workbook changed while reading. Refresh again.",
  "تغير Excel بعد فتح الشاشة. اضغط تحديث وراجع البيانات ثم احفظ.": "Excel changed after this screen opened. Refresh, review the data and save again.",
  "تغير الملف أثناء تجهيز الحفظ. حدّث المصدر.": "The workbook changed while preparing the save. Refresh the source.",
  "لم يصل تأكيد الحفظ. حدّث المصدر وتحقق من رقم الطلب قبل إعادة العملية.": "Save confirmation was not received. Refresh and check the request ID before retrying.",
  "لم يصل تأكيد نهائي. حدّث المصدر وتحقق من رقم الطلب.": "Final confirmation was not received. Refresh and check the request ID.",
  "mode يجب أن يكون local أو onedrive.": "Storage mode must be local or onedrive.",
  "أكمل إعدادات Microsoft في Secrets حسب دليل التثبيت.": "Complete the Microsoft Secrets settings as described in the installation guide.",
  "بداية الفترة بعد نهايتها.": "The start date is after the end date.",
  "أضف سجلًا نشطًا في ": "Add an active record in ",
  " أولًا.": " first.",
  " ريال.": " SAR."
})

EN_TEXT.update({"الأشهر": "Months", "الحد الائتماني": "Credit limit", "الدخل": "Income", "الدفعة الأولى": "Down payment", "القيمة": "Value", "المفتاح": "Key", "الوسيلة": "Contact method", "تاريخ الإلغاء": "Cancellation date", "سبب الإلغاء": "Cancellation reason", "آخر حفظ": "Last saved", "الموعد": "Due date"})

def tr(value,language=None):
    language=language if language is not None else st.session_state.get('ui_language','ar')
    if not isinstance(value,str) or language=='ar':return value
    if value in EN_TEXT:return EN_TEXT[value]
    if not re.search(r'[\u0600-\u06ff]',value):return value
    # Translate known message fragments, leaving customer-entered names intact.
    for source in sorted(EN_TEXT,key=len,reverse=True):
        if (len(source)>=8 or source in [' ريال.',' أولًا.']) and source in value:value=value.replace(source,EN_TEXT[source])
    return value

def display_option(value,formatter=None,language=None):
    return tr(formatter(value) if formatter else str(value),language)

THEME_COLORS = {
  "light": {
    "bg": "#b9cecb",
    "panel": "#acb7b3",
    "card": "#d4d9d6",
    "text": "#273b36",
    "muted": "#354a42",
    "border": "#93a6a0",
    "input_border": "#687f75",
    "accent": "#2b4c44",
    "button": "#315f56",
    "button_text": "#f7f8f4",
    "hover": "#bac6c0",
    "selection": "#a6bdb4",
    "positive": "#315f56",
    "warning": "#715522",
    "negative": "#874b3d",
    "neutral": "#445b52",
    "alerts": {
      "Info": [
        "#b4c9ce",
        "#314f59"
      ],
      "Success": [
        "#bacdbd",
        "#2b5039"
      ],
      "Warning": [
        "#d2c6aa",
        "#5b461d"
      ],
      "Error": [
        "#d5bbb0",
        "#713c2d"
      ]
    },
    "field": "#c6ceca"
  },
  "dark": {
    "bg": "#19211f",
    "panel": "#202b28",
    "card": "#26322e",
    "text": "#d2dbd5",
    "muted": "#a7b6ad",
    "border": "#42554b",
    "input_border": "#788f81",
    "accent": "#9abbb0",
    "button": "#365e50",
    "button_text": "#f0f4ee",
    "hover": "#33463c",
    "selection": "#3d5145",
    "positive": "#9abbb0",
    "warning": "#d3be91",
    "negative": "#dcaa9f",
    "neutral": "#a7b6ad",
    "alerts": {
      "Info": [
        "#28363a",
        "#b7c8cf"
      ],
      "Success": [
        "#2a3930",
        "#b0c8b4"
      ],
      "Warning": [
        "#3a3529",
        "#d4c29b"
      ],
      "Error": [
        "#3c2e2a",
        "#d9b0a2"
      ]
    },
    "field": "#26322e"
  }
}

def theme_colors():
    return THEME_COLORS['dark' if st.session_state.get('ui_theme')=='dark' else 'light']

def appearance():
    language=st.query_params.get('lang','ar')
    if language not in ['ar','en']:language='ar'
    st.session_state.setdefault('ui_language',language)
    st.session_state.setdefault('ui_theme','dark' if st.query_params.get('theme')=='dark' else 'light')
    st.sidebar.selectbox(tr('اللغة / Language'),['ar','en'],format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:'العربية / Arabic' if x=='ar' else 'English / الإنجليزية'),language),key='ui_language')
    st.query_params['lang']=st.session_state['ui_language']
    dark=st.session_state['ui_theme']=='dark'
    if st.sidebar.button(tr(('☀️ الوضع النهاري' if dark else '🌙 الوضع الليلي') if st.session_state['ui_language']=='ar' else ('☀️ Light mode' if dark else '🌙 Dark mode')),key='theme_switch',width='stretch'):
        st.session_state['ui_theme']='light' if dark else 'dark'
        st.query_params['theme']=st.session_state['ui_theme'];st.rerun()
    rtl=st.session_state['ui_language']=='ar'
    colors=theme_colors()
    bg,panel,card,text,muted,border,accent=(colors[k] for k in ['bg','panel','card','text','muted','border','accent'])
    direction='rtl' if rtl else 'ltr';align='right' if rtl else 'left'
    # Only reverse the layout row. Never set RTL on the app shell: Streamlit's
    # sidebar collapse/width calculation requires the shell to remain LTR.
    layout='row-reverse' if rtl else 'row'
    css=f"""<style>
    :root{{--app-bg:{bg};--app-panel:{panel};--app-card:{card};--app-text:{text};--app-muted:{muted};--app-border:{border};--app-accent:{accent};color-scheme:{'dark' if dark else 'light'};}}
    .stApp{{direction:ltr;color-scheme:{"dark" if dark else "light"};background:{bg};color:{text};font-family:Tahoma,Arial,sans-serif;}}
    [data-testid="stAppViewContainer"]{{flex-direction:{layout};background:{bg};}}
    [data-testid="stMain"]{{min-width:0;background:{bg};}}
    [data-testid="stSidebar"][data-collapsed="true"]{{display:none!important;visibility:hidden;pointer-events:none;transform:translateX({'100%' if rtl else '-100%'});}}
    [data-testid="stSidebarCollapseButton"] svg,[data-testid="stExpandSidebarButton"] svg{{transform:{'rotate(180deg)' if rtl else 'none'};}}
    [data-testid="stSidebarResizeHandle"]{{display:{'none' if rtl else 'block'};}}
    [data-testid="stHeader"]{{background:{bg};color:{text};}}
    [data-testid="stSidebar"]{{direction:ltr;background:{panel};color:{text};}}
    [data-testid="stSidebarUserContent"]{{direction:{direction};text-align:{align};padding-top:2rem;}}
    [data-testid="stMainBlockContainer"],.block-container{{direction:{direction};text-align:{align};max-width:1550px;padding-top:4.5rem;}}
    [data-testid="stMarkdownContainer"],[data-testid="stText"],[data-testid="stExpander"] summary{{color:{text};}}
    h1,h2,h3,p,label{{font-family:Tahoma,Arial,sans-serif!important;}}
    h1,h2,h3,[data-testid="stWidgetLabel"],[data-testid="stMetricLabel"],.stCaption{{color:{text}!important;text-align:{align};}}
    h1{{font-size:2rem!important;}} [data-testid="stCaptionContainer"]{{color:{muted};}}
    [data-testid="stMetric"]{{background:{card};border:1px solid {border};border-radius:14px;padding:16px;margin:5px 0;}}
    [data-testid="stMetricValue"]{{color:{accent};font-size:1.7rem;direction:ltr;text-align:{align};}}
    [data-testid="stForm"]{{border:1px solid {border};border-radius:12px;color:{text};}}
    [data-testid="stExpander"]{{background:{colors['field']};border:1px solid {border};border-radius:12px;color:{text};}}
    input,textarea,[data-baseweb="input"],[data-baseweb="base-input"],[data-baseweb="textarea"],[data-baseweb="select"]>div{{background:{colors['field']}!important;color:{text}!important;caret-color:{text};}}
    input,textarea{{direction:{direction};text-align:{align};}} input[type="number"]{{direction:ltr;}}
    [data-baseweb="select"] span,[data-baseweb="select"] svg{{color:{text};fill:{text};}}
    [data-baseweb="popover"] [role="listbox"],[data-baseweb="popover"] li,[data-baseweb="popover"] [data-baseweb="calendar"]{{background:{card}!important;color:{text}!important;}}
    [data-testid="stBaseButton-secondary"],[data-testid="stBaseButton-secondaryFormSubmit"],[data-testid="stBaseButton-tertiary"]{{background:{colors['field']};color:{text};border-color:{border};border-radius:9px;}}
    [data-testid="stBaseButton-primary"],[data-testid="stBaseButton-primaryFormSubmit"]{{background:{colors['button']};color:{colors['button_text']};border-color:transparent;border-radius:9px;}}
    input::placeholder,textarea::placeholder{{color:{muted}!important;opacity:1;}}
    [data-baseweb="input"],[data-baseweb="select"]>div,[data-baseweb="textarea"]{{border-color:{colors['input_border']}!important;}}
    [data-baseweb="input"]:focus-within,[data-baseweb="select"]:focus-within>div,[data-baseweb="textarea"]:focus-within{{border-color:{accent}!important;box-shadow:0 0 0 2px {accent}33!important;}}
    [data-testid="stBaseButton-secondary"]:hover,[data-testid="stBaseButton-secondaryFormSubmit"]:hover{{background:{colors['hover']};border-color:{colors['input_border']};color:{text};}}
    button:focus-visible{{outline:2px solid {accent};outline-offset:3px;}}
    button:disabled{{opacity:0.58;}}
    a{{color:{accent};}} ::selection{{background:{colors['selection']};color:{text};}}
    [data-testid="stAlert"]{{color:{text};}} [data-testid="stAlert"] p{{color:inherit;}}
    [data-testid="stDataFrame"] .dvn-scroller{{background:{card};--gdg-bg-cell:{card}!important;--gdg-bg-cell-medium:{panel}!important;--gdg-bg-header:{panel}!important;--gdg-bg-header-hovered:{card}!important;--gdg-bg-header-has-focus:{panel}!important;--gdg-text-dark:{text}!important;--gdg-text-medium:{muted}!important;--gdg-text-light:{muted}!important;--gdg-text-header:{text}!important;--gdg-border-color:{border}!important;}}
    </style>"""
    mobile_css=f"""
    @media (max-width: 768px) {{
      [data-testid="stMainBlockContainer"],.block-container{{padding:3.8rem max(0.85rem,env(safe-area-inset-right)) max(2rem,env(safe-area-inset-bottom)) max(0.85rem,env(safe-area-inset-left));max-width:100%;}}
      [data-testid="stSidebar"][data-collapsed="false"]{{position:absolute!important;top:0;bottom:0;{'right:0;left:auto;' if rtl else 'left:0;right:auto;'}width:min(88vw,21rem)!important;min-width:0!important;max-width:88vw!important;box-shadow:{'-8px' if rtl else '8px'} 0 24px #00000022;}}
      [data-testid="stSidebarUserContent"]{{padding-top:0.75rem;padding-bottom:max(1.5rem,env(safe-area-inset-bottom));}}
      [data-testid="stSidebar"] h1{{font-size:1.3rem!important;}}
      h1{{font-size:1.6rem!important;}} h2{{font-size:1.2rem!important;}} h3{{font-size:1.05rem!important;}}
      [data-testid="stHorizontalBlock"]{{gap:0.75rem;}}
      [data-testid="stHorizontalBlock"]:has([data-testid="stMetric"]){{display:grid;grid-template-columns:repeat(2,minmax(0,1fr));align-items:stretch;}}
      [data-testid="stHorizontalBlock"]:has([data-testid="stMetric"])>[data-testid="stColumn"]{{width:100%!important;min-width:0!important;flex:none!important;}}
      [data-testid="stHorizontalBlock"]:not(:has([data-testid="stMetric"])){{flex-direction:column;}}
      [data-testid="stHorizontalBlock"]:not(:has([data-testid="stMetric"]))>[data-testid="stColumn"]{{width:100%!important;min-width:0!important;}}
      [data-testid="stMetric"]{{padding:12px 10px;border-radius:10px;margin:0;}}
      [data-testid="stMetricValue"]{{font-size:1.4rem;}}
      [data-testid="stMetricLabel"]{{white-space:normal;overflow-wrap:anywhere;}}
      [data-testid="stForm"]{{padding:0.85rem;}}
      input,textarea,select,[data-baseweb="select"]{{font-size:16px!important;}}
      input,[data-baseweb="select"]>div,button{{min-height:44px;}}
      textarea{{min-height:90px;}}
      button{{white-space:normal;overflow-wrap:anywhere;touch-action:manipulation;}}
      [data-testid="stCheckbox"] label,[data-testid="stRadio"] label{{min-height:44px;align-items:center;}}
      [data-testid="stWidgetLabel"] p,[data-testid="stMarkdownContainer"] p{{overflow-wrap:anywhere;}}
      [data-testid="stDataFrame"]{{max-width:100%;min-width:0;}}
      [data-testid="stTabs"] [role="tablist"]{{overflow-x:auto;flex-wrap:nowrap;}}
      [data-testid="stTabs"] [role="tab"]{{min-height:44px;flex-shrink:0;}}
      [data-baseweb="popover"]{{max-width:calc(100vw - 16px);}}
      [data-testid="stCaptionContainer"],[data-testid="stCaptionContainer"] p{{color:{colors['muted']}!important;}}
    }}
    @media (max-width: 350px) {{
      [data-testid="stHorizontalBlock"]:has([data-testid="stMetric"]){{grid-template-columns:minmax(0,1fr);}}
    }}
    """
    css=css.replace('</style>',mobile_css+'</style>')
    for kind,(alert_bg,alert_text) in colors['alerts'].items():
        css=css.replace('</style>',f"""
        [data-testid="stAlertContainer"]:has([data-testid="stAlertContent{kind}"]){{background:{alert_bg};color:{alert_text};border:1px solid {border};border-radius:10px;}}
        [data-testid="stAlertContent{kind}"],[data-testid="stAlertContent{kind}"] [data-testid="stMarkdownContainer"],[data-testid="stAlertContent{kind}"] p{{color:{alert_text}!important;}}
        </style>""")
    st.markdown(css,unsafe_allow_html=True)

def field_help(label):
    hints={
      'الاسم':'الاسم الكامل كما يظهر في المستندات.',
      'الجوال':'إلزامي: 10 أرقام تبدأ بـ05، مثال 0512345678.',
      'الهوية':'10 أرقام بالضبط دون مسافات؛ إلزامي للعميل. للمستثمر اختياري حتى تتوفر هويته.',
      'تكلفة الشراء الفعلية':'المبلغ المدفوع فعليًا للمورّد بالريال؛ أساس حساب تكلفة الصفقة.',
      'سعر الكاش / أصل التمويل للعميل':'سعر البيع النقدي المرجعي؛ قد يختلف عن تكلفة شرائك الفعلية.',
      'سعر البيع الإجمالي شامل الربح والدفعة الأولى':'سعر البيع الإجمالي شامل الدفعة الأولى والربح، دون رسوم العميل الإضافية. لا تضف الربح إليه مرة أخرى.',
      'التمويل المحول لشراء السلعة (دون الرسوم)':'المبلغ المحول للشراء فقط؛ فائضه عن التكلفة يبقى حقًا لصاحب التمويل.',
      'رسوم مدفوعة للمكتب المنفذ':'رسوم تدفعها أنت للمكتب الخارجي؛ تزيد تكلفتك وتقلل ربحك.',
      'رسومك الثابتة على المستثمر':'رسم مستلم من المستثمر إضافة إلى تمويل الشراء؛ دخل للمكتب وحده.',
      'رسوم إضافية يدفعها العميل خارج الأقساط':'رسم مستقل للمكتب يُحصّل عند التنفيذ ولا يدخل ضمن سعر البيع أو سند لأمر تلقائيًا.',
      'مصاريف أخرى تتحملها أنت':'تكاليف الصفقة بالريال يتحملها المكتب، بما في ذلك العقود المدارة.',
      'عدد الأقساط الشهرية':'عدد الدفعات الشهرية بعد الدفعة الأولى؛ من 1 إلى 120.',
      'تاريخ العقد':'تاريخ إبرام العقد؛ لا يتجاوز اليوم.',
      'أول استحقاق':'موعد أول قسط شهري، لا يسبق تاريخ العقد.',
      'قيمة سند لأمر':'قيمة السند الفعلي بالريال؛ للتوثيق ولا تُضاف إلى المديونية. صفر يعني لم تُسجل.',
      'المبلغ المستلم فعليًا':'المبلغ الذي دفعه العميل فعلًا؛ لا يتجاوز المتبقي، ويوزع على أقدم الأقساط.',
      'خصم التسوية المتفق عليه (للتسوية المبكرة فقط)':'في التسوية: المبلغ المستلم + الخصم يجب أن يساويا كامل المتبقي.',
      'الحد الائتماني (صفر = غير محدد)':'الحد الأقصى للتعرض حسب أصل التمويل. صفر يلغي تطبيق الحد.',
      'المستوى':'تصنيف داخلي يدوي من A إلى D؛ لا يمثل فحصًا ائتمانيًا خارجيًا.',
      'الحالة':'نشط يسمح بعقود جديدة؛ موقوف يمنعها؛ مؤرشف للسجلات بلا عقود قائمة.',
      'مرجع السداد':'رقم العملية البنكية أو الإيصال؛ يجب ألا يتكرر مع عملية مالية أخرى.',
      'مرجع التحويل / العملية':'رقم العملية أو مرجع مستند الشراء لتمييزها ومنع التكرار.',
      'وعد السداد':'تاريخ التزام العميل بالسداد؛ يظهر التنبيه إذا تجاوزه دون إغلاق المتابعة.',
      'المتابعة القادمة':'الموعد الذي يجب التواصل فيه مع العميل مرة أخرى.',
      'موعد المتابعة القادمة':'يظهر في تنبيهات الملخص عند حلول الموعد إذا بقيت المتابعة مفتوحة.',
      'المبلغ':'المبلغ الفعلي بالريال؛ يجب أن يكون أكبر من صفر لحفظ حركة مالية.',
      'سبب إلغاء العقد':'سبب الفسخ أو تصحيح القيد. يحتفظ النظام به مع العقد الملغى.',
    }
    label=str(label)
    if label in hints:return hints[label]
    if 'تاريخ' in label or 'موعد' in label:return 'اختر التاريخ الصحيح للعملية أو المتابعة.'
    if 'ملاحظ' in label or 'وصف' in label or 'نتيجة' in label:return 'اكتب تفاصيل تساعدك على مراجعة السجل لاحقًا؛ لا تدخل معلومات سرية غير مطلوبة.'
    if 'مبلغ' in label or 'قيمة' in label or 'رسوم' in label or 'الدخل' in label:return 'أدخل القيمة الفعلية بالريال السعودي دون فواصل؛ تُحسب المبالغ بدقة الهللة.'
    if st.session_state.get('ui_language','ar')=='en':return 'Select '+tr(label)+' for this record; check your selection before saving.'
    return 'حدد '+label+' لهذا السجل؛ تحقق من الاختيار قبل الحفظ.'

def format_money(h):return f'{sar(h):,.2f}'
def show_table(rows):
    if not rows:st.info(tr('لا توجد سجلات حتى الآن.'));return
    clean=[{k:(v.isoformat() if isinstance(v,(date,datetime)) else v) for k,v in r.items()} for r in rows]
    if st.session_state.get('ui_language','ar')=='en':
        controlled={'الحالة','النوع','الوسيلة','البيان','المستوى'}
        clean=[{tr(k):(tr(v) if k in controlled else v) for k,v in row.items()} for row in clean]
    frame=pd.DataFrame(clean)
    if st.session_state.get('ui_language','ar')=='ar':frame=frame.iloc[:,::-1]
    dark=st.session_state.get('ui_theme')=='dark'
    colors=theme_colors()
    styled=frame.style.set_properties(**{'background-color':colors['card'],'color':colors['text']})
    if tr('الحالة') in frame.columns:
        status_colors={'متأخر':colors['negative'],'ملغى':colors['neutral'],'منتظم':colors['positive'],'مكتمل':colors['positive'],'مسدد':colors['positive'],'مسدد مبكرًا':colors['positive'],'مسودة':colors['warning'],'مستحق اليوم':colors['warning']}
        palette={tr(k):v for k,v in status_colors.items()};fallback=colors['text']
        styled=styled.map(lambda x:'color: '+palette.get(str(x),fallback)+'; font-weight: bold',subset=[tr('الحالة')])
    st.dataframe(styled,hide_index=True,width='stretch')
def metrics(items):
    cols=st.columns(min(4,len(items)))
    for i,(label,value) in enumerate(items):cols[i%len(cols)].metric(tr(label),value)
def names(data,sheet):return {r['المعرف']:r['الاسم'] for r in data[sheet]}
def choose(data,sheet,label,key,allow_empty=False,active=False,default=None):
    rows=[r for r in data[sheet] if not active or r['الحالة']=='نشط'];mapping={r['المعرف']:r['الاسم'] for r in rows}
    options=([''] if allow_empty else [])+list(mapping)
    if not options:st.info(tr('أضف سجلًا نشطًا في '+sheet+' أولًا.'));return None
    return st.selectbox(tr(label),options,index=options.index(default) if default in options else 0,format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:mapping.get(x,'بدون')),language),key=key,help=tr(field_help(label)))
def save(action,payload):
    req=st.session_state.setdefault('request',uid('REQ'))
    try:
        raw,tag=st.session_state['store'].commit(st.session_state['tag'],action,payload,req)
        st.session_state['raw']=raw;st.session_state['tag']=tag;st.session_state.pop('request',None)
        st.session_state['last_saved']=datetime.now(TZ).strftime('%Y-%m-%d %H:%M:%S');st.session_state['notice']='تم الحفظ في Excel بنجاح.';st.rerun()
    except Exception as exc:
        st.error(tr(str(exc)));st.caption(tr('رقم الطلب: '+req+'. عند انقطاع الاتصال، حدّث المصدر قبل إعادة الحفظ.'))

def contract_rows(data,ss):
    cl=names(data,'العملاء');iv=names(data,'المستثمرون');office=names(data,'المكاتب')
    return [{'العقد':s['id'],'العميل':cl.get(s['c']['العميل'],''),'النوع':s['c']['النوع'],'المستثمر':iv.get(s['c'].get('المستثمر'),''),'المكتب':office.get(s['c'].get('المكتب'),''),
      'نوع السلعة':details(data,s['id']).get('نوع السلعة',''),'الدفعة الأولى':sar(s['down']),'قيمة سند لأمر':sar(money(details(data,s['id']).get('قيمة سند لأمر'))),'تكلفة الشراء':sar(s['p']['purchase']),'سعر الكاش':sar(s['p']['cash']),'إجمالي التقسيط':sar(s['p']['total']),'ربح المكتب المتوقع':sar(s['p']['office_profit']),
      'المحصل':sar(s['paid']),'الخصم':sar(s['waived']),'المتبقي':sar(s['remaining']),'المتأخر':sar(s['overdue']),'الاستحقاق القادم':s['next_due'],'الحالة':s['status']} for s in ss]
def part_rows(pp):
    return [{'القسط':z['sequence'],'العقد':z['contract'],'الاستحقاق':z['due'],'المبلغ':sar(z['amount']),'المحصل':sar(z['paid']),'الخصم':sar(z['waived']),'المتبقي':sar(z['remaining']),'الحالة':z['status'],'أيام التأخير':z['late_days']} for z in pp]

def home(data,snap):
    funded=[s for s in snap['contracts'] if s['active']]
    st.subheader(tr('التنبيهات والمتابعة'))
    overdue=[s for s in funded if s['overdue']>0]
    follows=[r for r in data['المتابعات'] if r['الحالة']=='مفتوحة' and ((dt(r.get('المتابعة القادمة')) and dt(r['المتابعة القادمة'])<=today()) or (dt(r.get('وعد السداد')) and dt(r['وعد السداد'])<today()))]
    metrics([('عقود متأخرة',str(len(overdue))),('متابعات مستحقة',str(len(follows))),('إجمالي المتأخر',format_money(sum(s['overdue'] for s in funded)))])
    cl=names(data,'العملاء')
    targets={s['id']:s for s in overdue}
    for r in follows:
        for x in snap['contracts']:
            if x['c']['العميل']==r['العميل'] and (not r.get('العقد') or r['العقد']==x['id']):targets[x['id']]=x
    if targets:
        target=st.selectbox(tr('فتح عقد من التنبيهات'),list(targets),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:cl.get(targets[x]['c']['العميل'],'')+' · '+x),language),help=tr(field_help('فتح عقد من التنبيهات')))
        if st.button(tr('فتح العقد المحدد')):
            st.session_state['overview_contract']=target;st.session_state['overview_action']='عرض / تعديل عقد'
            st.session_state['overview_search']='';st.session_state['overview_status']='الكل';st.session_state['overview_investor']=''
    with st.expander(tr('المستحق والمتأخر ومتابعات اليوم')):
        show_table(part_rows([z for z in snap['parts'] if z['active'] and z['remaining']>0 and (z['due']-today()).days<=7]))
        show_table([{'العميل':cl.get(r['العميل'],''),'المتابعة':r.get('المتابعة القادمة'),'وعد السداد':r.get('وعد السداد'),'النتيجة':r.get('النتيجة')} for r in follows])
    st.subheader(tr('الملخص المالي'))
    metrics([('السيولة الحرة',format_money(snap['free'])),('رصيد البنك',format_money(snap['bank'])),('أمانات المستثمرين',format_money(snap['held'])),('ربح المكتب المتوقع',format_money(sum(s['p']['office_profit'] for s in funded))),('متبقي عقودك',format_money(sum(s['remaining'] for s in funded if s['c']['النوع']!=MANAGED))),('متبقي عقود المستثمرين',format_money(sum(s['remaining'] for s in funded if s['c']['النوع']==MANAGED)))])
    st.caption(tr('الربح المتوقع يفترض اكتمال السداد؛ أمانات المستثمرين مستبعدة من السيولة الحرة.'))
    st.subheader(tr('العقود والإجراءات السريعة'))
    quick=st.columns(3)
    for col,label in zip(quick,['إضافة عميل','عقد جديد','تسجيل سداد']):
        if col.button(tr(label),key='quick-'+label,width='stretch'):st.session_state['overview_action']=label
    action=st.selectbox(tr('الإجراء'),['عرض / تعديل عقد','إضافة عميل','عقد جديد','تسجيل سداد'],key='overview_action',help=tr(field_help('الإجراء')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
    if action=='إضافة عميل':people(data,snap,'العملاء');return
    if action=='عقد جديد':contract_form(data,key='overview-new');return
    query=st.text_input(tr('بحث بالعميل أو رقم العقد أو السلعة'),key='overview_search',help=tr(field_help('بحث بالعميل أو رقم العقد أو السلعة')))
    status=st.selectbox(tr('حالة العقد'),['الكل','مسودة','منتظم','متأخر','مكتمل','مسدد مبكرًا','ملغى'],key='overview_status',help=tr(field_help('حالة العقد')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
    investors=names(data,'المستثمرون');inv=st.selectbox(tr('تصفية بالمستثمر'),['']+list(investors),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:investors.get(x,'الكل')),language),key='overview_investor',help=tr(field_help('تصفية بالمستثمر')))
    selected=[s for s in snap['contracts'] if (status=='الكل' or s['status']==status) and (not inv or s['c'].get('المستثمر')==inv) and (not query or query in (cl.get(s['c']['العميل'],'')+' '+s['id']+' '+str(details(data,s['id']).get('نوع السلعة') or '')))]
    show_table(contract_rows(data,selected))
    if selected:
        mapping={s['id']:s for s in selected}
        if st.session_state.get('overview_contract') not in mapping:st.session_state.pop('overview_contract',None)
        cid=st.selectbox(tr('اختر العقد لعرضه أو تعديله'),list(mapping),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:cl.get(mapping[x]['c']['العميل'],'')+' · '+x),language),key='overview_contract',help=tr(field_help('اختر العقد لعرضه أو تعديله')))
        contract_panel(data,snap,mapping[cid])

def delete_control(sheet,ident):
    with st.expander(tr('حذف السجل نهائيًا'),expanded=False):
        st.caption(tr('الحذف للسجلات غير المرتبطة فقط. الحركات المالية تُصحح بالإلغاء للحفاظ على سجلها.'))
        confirm=st.checkbox(tr('أؤكد حذف هذا السجل'),key='delete-ok-'+sheet+ident,help=tr(field_help('أؤكد حذف هذا السجل')))
        if st.button(tr('حذف السجل'),disabled=not confirm,key='delete-'+sheet+ident):save('delete_record',{'sheet':sheet,'id':ident})

def people(data,snap,sheet):
    cl=sheet=='العملاء';inv=sheet=='المستثمرون'
    action='save_client' if cl else ('save_investor' if inv else 'save_office')
    mapping=names(data,sheet);ident=st.selectbox(tr('إضافة أو تعديل'),['']+list(mapping),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:mapping.get(x,'إضافة سجل جديد')),language),key='person-'+sheet,help=tr(field_help('إضافة أو تعديل')))
    old=lookup(data,sheet,ident) if ident else {};p={}
    with st.form('person-form-'+sheet+'-'+ident):
        a,b=st.columns(2)
        with a:
            p['الاسم']=st.text_input(tr('الاسم'),old.get('الاسم') or '',help=tr(field_help('الاسم')))
            p['الجوال']=st.text_input(tr('الجوال'),old.get('الجوال') or '',help=tr('إلزامي: 10 أرقام تبدأ بـ05، مثال 0512345678.'))
            if cl or inv:p['الهوية']=st.text_input(tr('الهوية'),old.get('الهوية') or '',help=tr('10 أرقام بالضبط؛ إلزامي للعميل. تُقبل الأرقام العربية والإنجليزية.'))
        with b:
            states=['نشط','موقوف','مؤرشف'];p['الحالة']=st.selectbox(tr('الحالة'),states,index=states.index(old.get('الحالة') or 'نشط'),help=tr(field_help('الحالة')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
            if cl:
                grades=['غير مصنف','A','B','C','D'];p['المستوى']=st.selectbox(tr('المستوى'),grades,index=grades.index(old.get('المستوى') or 'غير مصنف'),help=tr(field_help('المستوى')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
                p['الحد الائتماني']=st.number_input(tr('الحد الائتماني (صفر = غير محدد)'),min_value=0.0,value=float(old.get('الحد الائتماني') or 0),help=tr(field_help('الحد الائتماني (صفر = غير محدد)')))
        if cl:
            a,b=st.columns(2)
            with a:
                for k in ['العنوان','جهة العمل','الكفيل']:p[k]=st.text_input(tr(k),old.get(k) or '',help=tr(field_help(k)))
            with b:
                p['الدخل']=st.number_input(tr('الدخل الشهري'),min_value=0.0,value=float(old.get('الدخل') or 0),help=tr(field_help('الدخل الشهري')));p['جوال الكفيل']=st.text_input(tr('جوال الكفيل'),old.get('جوال الكفيل') or '',help=tr('اختياري؛ عند إدخاله يجب أن يكون 10 أرقام تبدأ بـ05.'))
        p['ملاحظات']=st.text_area(tr('ملاحظات'),old.get('ملاحظات') or '',help=tr(field_help('ملاحظات')))
        if st.form_submit_button(tr('حفظ في Excel'),type='primary'):save(action,{**p,**({'المعرف':ident} if ident else {})})
    if inv and ident:
        metrics([(label,format_money(value)) for label,value in investor_capital(data,snap,ident)])
        st.caption(tr('رأس المال الممول هو مجموع تمويل العقود المنفذة دون رسوم المكتب؛ رقم تراكمي قد يشمل إعادة تدوير المال، ولا يساوي الرصيد النقدي الحالي.'))
    if ident:delete_control(sheet,ident)
    if inv:
        show_table([{'الاسم':r['الاسم'],'الجوال':r['الجوال'],'رأس المال الممول':sar(dict(investor_capital(data,snap,r['المعرف']))['إجمالي رأس المال الممول']),'أموال محتجزة':sar(snap['holdings'].get(r['المعرف'],0)),'متبقي العقود':sar(sum(s['remaining'] for s in snap['contracts'] if s['active'] and s['c'].get('المستثمر')==r['المعرف'])),'الحالة':r['الحالة']} for r in data[sheet]])
    else:show_table(data[sheet])
    if ident:
        key='العميل' if cl else ('المستثمر' if inv else 'المكتب');show_table(contract_rows(data,[s for s in snap['contracts'] if s['c'].get(key)==ident]))

def contract_form(data,old=None,key='new',amend=False):
    old=old or {};p={};st.caption(tr('الحسابات تتحدث عند تعديل الحقول. حفظ العقد ينشئ مسودة؛ التنفيذ يسجل النقد والرسوم الفعلية.'))
    modes=MODES;p['النوع']=st.selectbox(tr('نوع الصفقة'),modes,index=modes.index(old.get('النوع',SELF)),key=key+'-mode',help=tr(field_help('نوع الصفقة')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
    p['العميل']=choose(data,'العملاء','العميل',key+'-client',active=not bool(old),default=old.get('العميل'))
    p['المستثمر']='';p['المكتب']=''
    if p['النوع']==MANAGED:
        p['المستثمر']=choose(data,'المستثمرون','المستثمر',key+'-investor',active=not bool(old),default=old.get('المستثمر'))
        st.info(tr('كامل تحصيلات الأقساط للمستثمر. رسومك الثابتة يدفعها المستثمر إضافة إلى مبلغ التمويل عند التنفيذ. فرق الشراء يبقى أمانة للمستثمر.'))
    if p['النوع']==EXTERNAL:p['المكتب']=choose(data,'المكاتب','المكتب المنفذ',key+'-office',active=not bool(old),default=old.get('المكتب'))
    a,b=st.columns(2)
    with a:
        p['تكلفة الشراء']=st.number_input(tr('تكلفة الشراء الفعلية'),min_value=0.0,value=float(old.get('تكلفة الشراء') or 4970),step=10.0,key=key+'-purchase',help=tr(field_help('تكلفة الشراء الفعلية')))
        p['سعر الكاش']=st.number_input(tr('سعر الكاش / أصل التمويل للعميل'),min_value=0.0,value=float(old.get('سعر الكاش') or 5000),step=100.0,key=key+'-cash',help=tr(field_help('سعر الكاش / أصل التمويل للعميل')))
        suggested=float(Decimal(str(p['سعر الكاش']))*(1+Decimal(str(setting(data,'default_profit_rate')))))
        p['إجمالي التقسيط']=st.number_input(tr('سعر البيع الإجمالي شامل الربح والدفعة الأولى'),min_value=0.0,value=float(old.get('إجمالي التقسيط') or suggested),step=100.0,key=key+'-total',help=tr(field_help('سعر البيع الإجمالي شامل الربح والدفعة الأولى')))
        st.caption(tr('القيمة المقترحة وفق الإعدادات: '+f'{suggested:,.2f}'+' ريال. عدّل إجمالي التقسيط حسب الاتفاق.'))
        p['التمويل المحول']=p['تكلفة الشراء'] if p['النوع']==SELF else st.number_input(tr('التمويل المحول لشراء السلعة (دون الرسوم)'),min_value=0.0,value=float(old.get('التمويل المحول') or 5000),step=100.0,key=key+'-fund',help=tr(field_help('التمويل المحول لشراء السلعة (دون الرسوم)')))
    with b:
        p['رسوم مكتب مدفوعة']=st.number_input(tr('رسوم مدفوعة للمكتب المنفذ'),min_value=0.0,value=float(old.get('رسوم مكتب مدفوعة') or 0),key=key+'-paidfee',help=tr(field_help('رسوم مدفوعة للمكتب المنفذ'))) if p['النوع']==EXTERNAL else 0
        p['رسوم مكتب ثابتة']=st.number_input(tr('رسومك الثابتة على المستثمر'),min_value=0.0,value=float(old.get('رسوم مكتب ثابتة') or 0),key=key+'-fee',help=tr(field_help('رسومك الثابتة على المستثمر'))) if p['النوع']==MANAGED else 0
        p['رسوم العميل']=st.number_input(tr('رسوم إضافية يدفعها العميل خارج الأقساط'),min_value=0.0,value=float(old.get('رسوم العميل') or 0),key=key+'-clientfee',help=tr(field_help('رسوم إضافية يدفعها العميل خارج الأقساط')))
        p['مصاريف أخرى']=st.number_input(tr('مصاريف أخرى تتحملها أنت'),min_value=0.0,value=float(old.get('مصاريف أخرى') or 0),key=key+'-expenses',help=tr(field_help('مصاريف أخرى تتحملها أنت')))
        p['الأشهر']=st.number_input(tr('عدد الأقساط الشهرية'),min_value=1,max_value=120,value=int(old.get('الأشهر') or setting(data,'default_months')),key=key+'-months',help=tr(field_help('عدد الأقساط الشهرية')))
        p['تاريخ العقد']=st.date_input(tr('تاريخ العقد'),value=dt(old.get('تاريخ العقد')) or today(),max_value=today(),key=key+'-signed',help=tr(field_help('تاريخ العقد')))
        p['أول استحقاق']=st.date_input(tr('أول استحقاق'),value=dt(old.get('أول استحقاق')) or month(today(),1),key=key+'-first',help=tr(field_help('أول استحقاق')))
    oldmeta=details(data,old.get('المعرف',''));extra={}
    extra['نوع السلعة']=st.text_input(tr('نوع السلعة'),oldmeta.get('نوع السلعة') or '',key=key+'-asset',help=tr('مثال: هاتف، أجهزة منزلية، سيارة. مطلوب للعقود الجديدة.'))
    extra['وصف السلعة']=st.text_input(tr('وصف السلعة'),oldmeta.get('وصف السلعة') or '',key=key+'-description',help=tr('الموديل أو الرقم التسلسلي أو وصف يميز السلعة.'))
    hasdown=st.checkbox(tr('هل توجد دفعة أولى؟'),value=money(oldmeta.get('الدفعة الأولى'))>0,key=key+'-hasdown',help=tr('جزء من سعر البيع الإجمالي يدفعه العميل عند التنفيذ، وليس رسمًا إضافيًا.'))
    extra['الدفعة الأولى']=st.number_input(tr('قيمة الدفعة الأولى'),min_value=0.0,value=float(oldmeta.get('الدفعة الأولى') or 0),key=key+'-down',help=tr('تُحصّل عند تنفيذ العقد وتُطرح من المبلغ المجدول على الأقساط.')) if hasdown else 0
    extra['قيمة سند لأمر']=st.number_input(tr('قيمة سند لأمر'),min_value=0.0,value=float(oldmeta.get('قيمة سند لأمر') or 0),key=key+'-note',help=tr('القيمة المثبتة في السند الفعلي؛ لا تُضاف إلى سعر البيع ولا إلى أرباح العقد. صفر يعني غير مسجل.'))
    p['ملاحظات']=st.text_area(tr('ملاحظات العقد'),old.get('ملاحظات') or '',key=key+'-notes',help=tr(field_help('ملاحظات العقد')))
    q=pricing(p);metrics([('فرق الشراء',format_money(q['saving'])),('ربح المكتب المتوقع',format_money(q['office_profit'])),('ربح المستثمر المتوقع',format_money(q['investor_profit'])),('القسط التقريبي',format_money(max(0,q['total']-money(extra['الدفعة الأولى']))//int(p['الأشهر'])))])
    if p['النوع']==EXTERNAL and q['surplus']>0:st.caption(tr('فائض لدى المكتب يجب استرداده وتسجيله: '+format_money(q['surplus'])+' ريال. الربح المتوقع يفترض استرداده.'))
    if p['النوع']==MANAGED:st.caption(tr('المطلوب استلامه من المستثمر عند التنفيذ: '+format_money(q['fund']+q['fee'])+' ريال. أمانته المتبقية بعد الشراء: '+format_money(q['surplus'])+' ريال.'))
    approved=st.checkbox(tr('راجعت شروط العقد الجديدة وهي تصحيح للصفقة والتحويلات الفعلية'),key=key+'-approve',help=tr(field_help('راجعت شروط العقد الجديدة وهي تصحيح للصفقة والتحويلات الفعلية'))) if amend else True
    if st.button(tr('حفظ تصحيح شروط العقد' if amend else 'حفظ مسودة العقد'),type='primary',key=key+'-save',disabled=not approved):
        if not p['العميل'] or (p['النوع']==MANAGED and not p['المستثمر']) or (p['النوع']==EXTERNAL and not p['المكتب']):st.error(tr('أضف واختر العميل والجهة الممولة أولًا.'))
        elif not txt(extra['نوع السلعة']):st.error(tr('أدخل نوع السلعة.'))
        elif hasdown and extra['الدفعة الأولى']<=0:st.error(tr('أدخل دفعة أولى أكبر من صفر أو ألغِ اختيار وجودها.'))
        else:save('amend_contract' if amend else 'save_contract',{**p,'details':extra,**({'المعرف':old['المعرف']} if old else {})})

def contracts_page(data,snap):
    new,existing=st.tabs([tr(v) for v in ['عقد جديد','العقود والتحصيل']])
    with new:contract_form(data)
    with existing:
        cl=names(data,'العملاء');ss={s['id']:s for s in snap['contracts']}
        show_table(contract_rows(data,snap['contracts']))
        if not ss:return
        cid=st.selectbox(tr('اختر العقد'),list(ss),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:cl.get(ss[x]['c']['العميل'],'')+' · '+x),language),help=tr(field_help('اختر العقد')))
        contract_panel(data,snap,ss[cid])

def contract_panel(data,snap,s):
    c=s['c'];p=s['p'];cid=s['id']
    if s['cancelled']:
        st.warning(tr('عقد ملغى: '+str(details(data,cid).get('سبب الإلغاء') or '')))
        delete_control('العقود',cid);return
    metrics([('إجمالي التقسيط',format_money(p['total'])),('المحصل',format_money(s['paid'])),('المتبقي',format_money(s['remaining'])),('الحالة',s['status'])])
    if not s['active']:
        st.write(tr('قيود التنفيذ التي ستسجل:'))
        show_table([{'البيان':r[0],'المبلغ':sar(r[1]),'أثر البنك':sar(r[2]),'أثر أمانات المستثمر':sar(r[3])} for r in activation_rows(c,details(data,cid))])
        with st.form('activate-'+cid):
            d=st.date_input(tr('تاريخ التنفيذ'),value=today(),max_value=today(),help=tr(field_help('تاريخ التنفيذ')));ref=st.text_input(tr('مرجع التحويل / العملية'),help=tr(field_help('مرجع التحويل / العملية')));notes=st.text_input(tr('ملاحظات التنفيذ'),help=tr(field_help('ملاحظات التنفيذ')))
            ok=st.checkbox(tr('تمت جميع التحويلات والشراء وتحصيل الرسوم الموضحة أعلاه فعليًا'),help=tr(field_help('تمت جميع التحويلات والشراء وتحصيل الرسوم الموضحة أعلاه فعليًا')))
            if st.form_submit_button(tr('تنفيذ الصفقة'),type='primary'):
                if not ok:st.error(tr('أكد إتمام العمليات الفعلية أولًا.'))
                else:save('activate',{'contract':cid,'date':d,'ref':ref,'notes':notes})
        with st.expander(tr('تعديل المسودة')):contract_form(data,c,key='edit-'+cid)
    elif s['remaining']>0:
        if c['النوع']==MANAGED:st.info(tr('كامل مبلغ السداد يضاف إلى أموال المستثمر المحتجزة.'))
        with st.form('receipt-'+cid):
            action=st.selectbox(tr('العملية'),['سداد قسط أو مبلغ جزئي','تسوية مبكرة'],help=tr(field_help('العملية')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
            amount=st.number_input(tr('المبلغ المستلم فعليًا'),min_value=0.0,max_value=sar(s['remaining']),value=min(sar(s['remaining']),sar((p['total']-s['down'])//int(c['الأشهر']))),help=tr(field_help('المبلغ المستلم فعليًا')))
            discount=st.number_input(tr('خصم التسوية المتفق عليه (للتسوية المبكرة فقط)'),min_value=0.0,value=0.0,help=tr(field_help('خصم التسوية المتفق عليه (للتسوية المبكرة فقط)')))
            d=st.date_input(tr('تاريخ السداد'),max_value=today(),help=tr(field_help('تاريخ السداد')));ref=st.text_input(tr('مرجع السداد'),help=tr(field_help('مرجع السداد')));notes=st.text_input(tr('المحصل / ملاحظات'),help=tr(field_help('المحصل / ملاحظات')))
            if st.form_submit_button(tr('تسجيل السداد'),type='primary'):save('settle' if action=='تسوية مبكرة' else 'receipt',{'contract':cid,'amount':amount,'discount':discount,'date':d,'ref':ref,'notes':notes})
    if s['surplus_due']>0:
        with st.form('refund-'+cid):
            st.write(tr('فائض مستحق لك لدى المكتب: '+format_money(s['surplus_due'])))
            amount=st.number_input(tr('المبلغ المسترد'),min_value=0.0,max_value=sar(s['surplus_due']),value=sar(s['surplus_due']),help=tr(field_help('المبلغ المسترد')));d=st.date_input(tr('تاريخ استرداد الفائض'),max_value=today(),help=tr(field_help('تاريخ استرداد الفائض')));ref=st.text_input(tr('مرجع الاسترداد'),help=tr(field_help('مرجع الاسترداد')))
            if st.form_submit_button(tr('تسجيل استرداد الفائض')):save('refund',{'contract':cid,'amount':amount,'date':d,'ref':ref})
    show_table(part_rows(s['parts']))
    meta=details(data,cid)
    st.caption(tr('السلعة: '+str(meta.get('نوع السلعة') or 'غير مسجل')+' · سعر الشراء: '+format_money(p['purchase'])+' · سعر البيع الإجمالي: '+format_money(p['total'])+' · الدفعة الأولى: '+format_money(s['down'])+' · سند لأمر: '+format_money(money(meta.get('قيمة سند لأمر')))))
    if s['active'] and not s['settled']:
        with st.expander(tr('تعديل الأسعار والدفعة الأولى والجدولة')):
            st.caption(tr('يحفظ التنفيذ السابق كملغى ويعيد حساب جميع الحركات. يرفض التصحيح إذا تعارض مع التحصيلات أو حقوق المستثمرين. لا يغيّر العقد الموقع خارج التطبيق.'))
            contract_form(data,c,key='amend-'+cid,amend=True)
    if s['active']:
        with st.expander(tr('تعديل بيانات السلعة والسند والملاحظات')):
            with st.form('info-'+cid):
                asset=st.text_input(tr('نوع السلعة'),meta.get('نوع السلعة') or '',help=tr(field_help('نوع السلعة')))
                description=st.text_input(tr('وصف السلعة'),meta.get('وصف السلعة') or '',help=tr(field_help('وصف السلعة')))
                note=st.number_input(tr('قيمة سند لأمر'),min_value=0.0,value=float(meta.get('قيمة سند لأمر') or 0),help=tr(field_help('قيمة سند لأمر')))
                notes=st.text_area(tr('ملاحظات العقد'),c.get('ملاحظات') or '',help=tr(field_help('ملاحظات العقد')))
                st.caption(tr('يمكن تعديل الأسعار والجدولة من قسم تصحيح شروط العقد أعلاه للعقود غير المسواة.'))
                if st.form_submit_button(tr('حفظ تعديلات العقد')):save('save_contract_info',{'contract':cid,'نوع السلعة':asset,'وصف السلعة':description,'قيمة سند لأمر':note,'ملاحظات':notes})
    with st.expander(tr('إلغاء العقد')):
        with st.form('cancel-'+cid):
            reason=st.text_input(tr('سبب إلغاء العقد'),help=tr(field_help('سبب إلغاء العقد')))
            reversed_funds=st.checkbox(tr('أؤكد إعادة السلعة وتسوية وإرجاع جميع مبالغ هذا العقد فعليًا، وأريد إلغاء قيوده'),help=tr(field_help('أؤكد إعادة السلعة وتسوية وإرجاع جميع مبالغ هذا العقد فعليًا، وأريد إلغاء قيوده')))
            st.caption(tr('الإلغاء يحتفظ بالعقد وسجله ويُلغي حركاته محاسبيًا؛ لا ينفذ تحويلات بنكية. يرفض إذا نتج رصيد سالب أو تعارض مع حقوق المستثمرين.'))
            if st.form_submit_button(tr('تأكيد إلغاء العقد')):save('cancel_contract',{'contract':cid,'reason':reason,'reversed':reversed_funds})
    if not s['active']:delete_control('العقود',cid)

def cash_page(data,snap):
    metrics([('البنك',format_money(snap['bank'])),('أمانات المستثمرين',format_money(snap['held'])),('المتاح للمكتب',format_money(snap['free']))])
    kind=st.selectbox(tr('نوع الحركة'),['إضافة رأس مال','سحب رأس مال','مصروف مكتب','تحويل للمستثمر'],help=tr(field_help('نوع الحركة')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
    inv=choose(data,'المستثمرون','المستثمر','transfer-investor') if kind=='تحويل للمستثمر' else ''
    if inv:st.caption(tr('المتاح تحويله للمستثمر: '+format_money(snap['holdings'].get(inv,0))))
    with st.form('cash'):
        amount=st.number_input(tr('المبلغ'),min_value=0.0,value=0.0,help=tr(field_help('المبلغ')));d=st.date_input(tr('تاريخ الحركة'),max_value=today(),help=tr(field_help('تاريخ الحركة')));ref=st.text_input(tr('المرجع'),help=tr(field_help('المرجع')));notes=st.text_input(tr('الوصف'),help=tr(field_help('الوصف')))
        if st.form_submit_button(tr('حفظ الحركة'),type='primary'):save('cash',{'kind':kind,'amount':amount,'date':d,'investor':inv,'ref':ref,'notes':notes})
    st.subheader(tr('سجل الحركات'));show_table(list(reversed(data['الحركات'])))
    with st.expander(tr('تعديل حركة مالية')):
        editable={t['المجموعة']:t for t in data['الحركات'] if t['الحالة']=='نشطة' and t['النوع'] in ['سداد','تسوية','استرداد فائض','إضافة رأس مال','سحب رأس مال','مصروف مكتب','تحويل للمستثمر']}
        if editable:
            gid=st.selectbox(tr('اختر حركة للتعديل'),list(editable),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:editable[x]['النوع']+' · '+str(editable[x]['المبلغ'])+' · '+x),language),help=tr(field_help('اختر حركة للتعديل')))
            t=editable[gid]
            with st.form('edit-tx-'+gid):
                amount=st.number_input(tr('المبلغ المصحح'),min_value=0.0,value=float(t['المبلغ']),help=tr(field_help('المبلغ المصحح')))
                d=st.date_input(tr('تاريخ الحركة المصحح'),value=dt(t['التاريخ']),max_value=today(),help=tr(field_help('تاريخ الحركة المصحح')))
                ref=st.text_input(tr('المرجع المصحح'),t.get('المرجع') or '',help=tr(field_help('المرجع المصحح')))
                notes=st.text_input(tr('وصف التصحيح'),t.get('ملاحظات') or '',help=tr(field_help('وصف التصحيح')))
                discount=st.number_input(tr('خصم التسوية المصحح'),min_value=0.0,value=float(t.get('خصم التسوية') or 0),help=tr(field_help('خصم التسوية المصحح'))) if t['النوع']=='تسوية' else 0
                if st.form_submit_button(tr('حفظ تصحيح الحركة')):save('edit_transaction',{'group':gid,'amount':amount,'date':d,'ref':ref,'notes':notes,'discount':discount})
    with st.expander(tr('تصحيح حركة بإلغاء المجموعة أو استعادتها')):
        groups={t['المجموعة']:t for t in data['الحركات']}
        if groups:
            group=st.selectbox(tr('المجموعة'),list(groups),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,(lambda x:x+' · '+str(groups[x]['التاريخ'])+' · '+groups[x]['الحالة']),language),help=tr(field_help('المجموعة')))
            ok=st.checkbox(tr('راجعت المجموعة كاملة وأريد تغيير حالتها'),help=tr(field_help('راجعت المجموعة كاملة وأريد تغيير حالتها')))
            if st.button(tr('إلغاء المجموعة' if groups[group]['الحالة']=='نشطة' else 'استعادة المجموعة'),disabled=not ok):save('void' if groups[group]['الحالة']=='نشطة' else 'restore',{'group':group})
        st.caption(tr('إلغاء التنفيذ يلغي قيوده معًا. إذا ترتب عليه سداد أو تحويل لاحق غير صالح فسيرفض النظام الإلغاء حتى تصحح الحركات التابعة.'))

def followups(data,snap):
    rows={r['المعرف']:r for r in data['المتابعات']};ident=st.selectbox(tr('إضافة أو تعديل متابعة'),['']+list(rows),help=tr(field_help('إضافة أو تعديل متابعة')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language));old=rows.get(ident,{})
    client=choose(data,'العملاء','العميل','follow-client')
    if old:client=old['العميل'];st.caption(tr('عميل المتابعة: '+names(data,'العملاء').get(client,'')))
    if not client:return
    options=['']+[c['المعرف'] for c in data['العقود'] if c['العميل']==client]
    with st.form('follow-'+ident):
        p={'العميل':client}
        p['العقد']=st.selectbox(tr('العقد (اختياري)'),options,index=options.index(old.get('العقد') or '') if (old.get('العقد') or '') in options else 0,help=tr(field_help('العقد (اختياري)')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
        p['تاريخ التواصل']=st.date_input(tr('تاريخ التواصل'),value=dt(old.get('تاريخ التواصل')) or today(),help=tr(field_help('تاريخ التواصل')))
        p['الوسيلة']=st.selectbox(tr('وسيلة التواصل'),['اتصال','واتساب','زيارة','رسالة','أخرى'],help=tr(field_help('وسيلة التواصل')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language));p['النتيجة']=st.text_input(tr('نتيجة التواصل'),old.get('النتيجة') or '',help=tr(field_help('نتيجة التواصل')))
        promise=st.checkbox(tr('يوجد وعد سداد'),value=bool(old.get('وعد السداد')),help=tr(field_help('يوجد وعد سداد')))
        pdue=st.date_input(tr('موعد الوعد'),value=dt(old.get('وعد السداد')) or today(),help=tr(field_help('موعد الوعد')));p['وعد السداد']=pdue if promise else None
        p['المبلغ الموعود']=st.number_input(tr('المبلغ الموعود'),min_value=0.0,value=float(old.get('المبلغ الموعود') or 0),help=tr(field_help('المبلغ الموعود')))
        p['المتابعة القادمة']=st.date_input(tr('موعد المتابعة القادمة'),value=dt(old.get('المتابعة القادمة')) or today(),help=tr(field_help('موعد المتابعة القادمة')));p['الحالة']=st.selectbox(tr('الحالة'),['مفتوحة','مكتملة'],index=1 if old.get('الحالة')=='مكتملة' else 0,help=tr(field_help('الحالة')),format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language))
        p['ملاحظات']=st.text_area(tr('الملاحظات'),old.get('ملاحظات') or '',help=tr(field_help('الملاحظات')))
        if st.form_submit_button(tr('حفظ المتابعة'),type='primary'):save('save_followup',{**p,**({'المعرف':ident} if ident else {})})
    if ident:delete_control('المتابعات',ident)
    show_table(data['المتابعات'])

def reports(data,snap):
    start=st.date_input(tr('من تاريخ'),value=today().replace(day=1),help=tr(field_help('من تاريخ')));end=st.date_input(tr('إلى تاريخ'),value=today(),help=tr(field_help('إلى تاريخ')))
    if start>end:st.error(tr('بداية الفترة بعد نهايتها.'));return
    txs=[t for t in snap['ledger'] if start<=dt(t['التاريخ'])<=end]
    bycontract={s['id']:s for s in snap['contracts']}
    managed=sum(money(t['المبلغ']) for t in txs if t['النوع'] in ['سداد','تسوية'] and bycontract[t['العقد']]['c']['النوع']==MANAGED)
    own=sum(money(t['المبلغ']) for t in txs if t['النوع'] in ['سداد','تسوية'] and bycontract[t['العقد']]['c']['النوع']!=MANAGED)
    fees=sum(money(t['المبلغ']) for t in txs if t['النوع'] in ['رسوم ثابتة','رسوم عميل'])
    metrics([('تحصيل أقساطك خلال الفترة',format_money(own)),('تحصيل للمستثمرين',format_money(managed)),('رسوم مكتسبة',format_money(fees)),('صافي أثر البنك',format_money(sum(money(t['أثر البنك']) for t in txs)))])
    st.caption(tr('تحصيل أقساطك يشمل رأس المال والربح. الرسوم المكتسبة منفصلة. صافي أثر البنك يشمل التمويل ورأس المال والمصروفات والتحويلات.'))
    show_table([{k:t.get(k) for k in HEADERS['الحركات'][:15]} for t in txs])
    st.subheader(tr('قائمة التعثر الحالية'));show_table(contract_rows(data,[s for s in snap['contracts'] if s['overdue']>0]))
    st.subheader(tr('حقوق المستثمرين الحالية'))
    show_table([{'المستثمر':r['الاسم'],'أمانات نقدية':sar(snap['holdings'][r['المعرف']]),'أقساط غير محصلة':sar(sum(s['remaining'] for s in snap['contracts'] if s['active'] and s['c'].get('المستثمر')==r['المعرف']))} for r in data['المستثمرون']])

def settings_page(data,snap):
    st.download_button(tr('تنزيل نسخة Excel الحالية'),st.session_state['raw'],file_name='Installment_System_V2.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary')
    with st.form('settings'):
        months=st.number_input(tr('عدد الأقساط الافتراضي'),min_value=1,max_value=120,value=int(setting(data,'default_months')),help=tr(field_help('عدد الأقساط الافتراضي')))
        rate=st.number_input(tr('الربح الإجمالي المقترح %'),min_value=0.0,max_value=1000.0,value=float(setting(data,'default_profit_rate'))*100,help=tr(field_help('الربح الإجمالي المقترح %')))
        if st.form_submit_button(tr('حفظ الإعدادات')):save('settings',{'default_months':months,'default_profit_rate':rate/100})
    st.write(tr('التصنيف A–D يدوي. فرق الشراء = سعر الكاش − التكلفة الفعلية. الربح المتوقع يحسب مرة واحدة من إجمالي التقسيط ناقص التكلفة والرسوم والمصاريف. الربح لا يضاف إلى سعر التقسيط مرة أخرى.'))
    st.write(tr('في العقد المدار: المستثمر يملك كامل الأقساط. رسومه الثابتة تدفع إضافة إلى التمويل وتتحول لملكية المكتب عند التنفيذ. فائض التمويل بعد الشراء يبقى أمانة للمستثمر حتى تحويله.'))
    st.write(tr('تعديل Excel خارج التطبيق متاح في حقول البيانات الأصلية وفق العناوين. الأقساط والملخص والحقول المحسوبة يعاد توليدها عند الحفظ. لتغيير الشروط المالية لعقد منفذ، صحح الحركات التابعة ثم ألغِ مجموعة التنفيذ أولًا.'))
    st.caption(tr('التطبيق يبدأ برصيد صفر. أضف رأس المال الحقيقي من صفحة الخزينة. ملف V1 مستقل ولا يُستبدل بهذا الملف.'))

def main():
    st.set_page_config(page_title='Installment Management | Ibrahim' if st.query_params.get('lang')=='en' else 'إدارة التقسيط | إبراهيم',page_icon='📊',layout='wide')
    appearance()
    st.sidebar.title(tr('إدارة التقسيط'));st.sidebar.caption(tr('إبراهيم · الإصدار '+VERSION))
    mode=os.getenv('INSTALLMENT_STORAGE_MODE',str(config('storage','mode','local')))
    try:
        if mode=='onedrive':
            import microsoft_auth as auth
            if not auth.cloud_mode_enabled():st.info(tr('أكمل إعدادات Microsoft في Secrets حسب دليل التثبيت.'));st.stop()
            token=auth.acquire_microsoft_token()
            if not token:st.stop()
            cfg=auth.microsoft_config();store=CloudStore(token,cfg['onedrive_folder'],cfg['workbook_name']);identity='cloud:'+store.path
            st.sidebar.success(tr('متصل بملف OneDrive V2'))
            if st.sidebar.button(tr('تسجيل الخروج')):
                st.session_state.clear();st.rerun()
        elif mode=='local':
            store=LocalStore(Path(os.getenv('INSTALLMENT_WORKBOOK',str(ROOT/'Installment_System_V2.xlsx'))));identity='local:'+str(store.path);st.sidebar.info(tr('الحفظ المحلي في Excel'))
        else:raise Error('mode يجب أن يكون local أو onedrive.')
        st.session_state['store']=store
        if 'raw' not in st.session_state or st.session_state.get('identity')!=identity:
            raw,tag=store.read();st.session_state['raw']=raw;st.session_state['tag']=tag;st.session_state['identity']=identity
        if st.sidebar.button(tr('↻ تحديث من Excel'),width='stretch'):
            raw,tag=store.read();st.session_state['raw']=raw;st.session_state['tag']=tag;st.rerun()
        data=read_excel(st.session_state['raw']);snap=snapshot(data)
    except Exception as exc:st.error(tr('تعذر فتح النظام: '+str(exc)));st.stop()
    st.sidebar.caption(tr('آخر حفظ بهذه الجلسة: '+st.session_state.get('last_saved','لم يُحفظ بعد')))
    pages=['الملخص','العملاء','المستثمرون','المكاتب','العقود والتحصيل','الخزينة والحركات','متابعة العملاء','التقارير','الإعدادات والنسخ']
    page=st.sidebar.radio(tr('الصفحات'),pages,format_func=lambda v, language=st.session_state.get('ui_language','ar'): display_option(v,language=language));st.title(tr(page))
    if 'notice' in st.session_state:st.success(tr(st.session_state.pop('notice')))
    try:
        if page in ['العملاء','المستثمرون','المكاتب']:people(data,snap,page)
        else:{'الملخص':home,'العقود والتحصيل':contracts_page,'الخزينة والحركات':cash_page,'متابعة العملاء':followups,'التقارير':reports,'الإعدادات والنسخ':settings_page}[page](data,snap)
    except Error as exc:st.error(tr(str(exc)))

if __name__=='__main__':main()
